# -*- coding: utf-8 -*-
"""Registro de Asistencia V12."""
import pickle
SP='/tmp/claude-0/-home-user-uniminuto/b364fe32-f339-5501-a525-09b94e80e92f/scratchpad'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.fill import PatternFillProperties, ColorChoice

D = pickle.load(open(SP+'/data.pkl','rb'))
FIRST, NROWS = 14, 2000
LAST = FIRST + NROWS - 1
EST_LAST, COL_LAST = 60001, 6001
EVA_FIRST, EVA_LAST = 10, 509
SEDES = ['BUG','BVA','CHI','CIN','FLO','IPI','PAS','PER','QUI']
TIPOS = ['ESTUDIANTE','PROFESOR','ADMINISTRATIVO','EXTERNO']
SLOTS, SHOW = 30, 15          # buscador

AZUL='1F3864'; AZUL2='2E5B9A'; CLARO='EAF1FA'; AMAR='FFF2CC'; GRIS='F7F7F7'; VERDE='E2EFDA'
thin=Side(style='thin',color='BFBFBF'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def Fo(sz=10,b=False,c='000000',i=False): return Font(name='Calibri',size=sz,bold=b,color=c,italic=i)
def Fi(c): return PatternFill('solid',fgColor=c)
def Al(h='center',v='center',w=False): return Alignment(horizontal=h,vertical=v,wrap_text=w)
def style(ws,ref,font=None,fill=None,align=None,border=None,fmt=None):
    for row in ws[ref]:
        for c in row:
            if font:c.font=font
            if fill:c.fill=fill
            if align:c.alignment=align
            if border:c.border=border
            if fmt:c.number_format=fmt
def band(ws,ref,text,bg=AZUL,fg='FFFFFF',size=11,h='left'):
    ws.merge_cells(ref); ws[ref.split(':')[0]]=text
    style(ws,ref,font=Fo(size,True,fg),fill=Fi(bg),align=Al(h,'center'))
def lbl(ws,ref,text):
    ws.merge_cells(ref); ws[ref.split(':')[0]]=text
    style(ws,ref,font=Fo(10,True,AZUL),align=Al('right','center'))
def dv(ws,f1,cells,tipo='list',**kw):
    v=DataValidation(type=tipo,formula1=f1,allow_blank=True,**kw)
    ws.add_data_validation(v)
    for c in cells: v.add(c)
    return v

wb=Workbook(); wb.remove(wb.active)

# ============================ REGISTROS ============================
rg=wb.create_sheet('Registros'); rg.sheet_properties.tabColor=AZUL
rg.sheet_view.showGridLines=False
ANCHOS={'A':7,'B':11,'C':15,'D':8,'E':34,'F':30,'G':32,'H':28,'I':13,'J':13,'K':17,
        'L':28,'M':17,'N':9,'O':26,'P':24,'Q':24,'R':2}
for k,v in ANCHOS.items(): rg.column_dimensions[k].width=v
for i in range(19,39):                       # S..AL auxiliares
    L=CL(i); rg.column_dimensions[L].width=12; rg.column_dimensions[L].hidden=True

band(rg,'A1:Q1','REGISTRO DE ASISTENCIA — BIENESTAR INSTITUCIONAL',AZUL,'FFFFFF',14)
rg.row_dimensions[1].height=26
band(rg,'A2:Q2','Escriba la cédula o el ID en «Documento» y el archivo trae los datos del participante.',
     AZUL2,'FFFFFF',10)

# ---- buscador tipo filtro ----
band(rg,'B4:G4','BUSCAR PARTICIPANTE POR NOMBRE',AZUL2,'FFFFFF',11)
lbl(rg,'B5:C5','Escriba parte del nombre:')
rg.merge_cells('D5:E5'); style(rg,'D5:E5',font=Fo(11,True,'7F6000'),fill=Fi(AMAR),align=Al('left'),border=BOX)
rg.merge_cells('F5:G5'); rg['F5']='Estudiantes'
style(rg,'F5:G5',font=Fo(10,True,'7F6000'),fill=Fi(AMAR),align=Al('center'),border=BOX)
dv(rg,'"Estudiantes,Colaboradores"',['F5'])
lbl(rg,'B6:C6','Seleccione el nombre:')
rg.merge_cells('D6:G6'); style(rg,'D6:G6',font=Fo(11,True,'7F6000'),fill=Fi(AMAR),align=Al('left'),border=BOX)
dv(rg,'=$AJ$5:$AJ$%d'%(4+SHOW),['D6'])
for r,t in ((7,'Documento:'),(8,'Sede y programa:'),(9,'Correo institucional:')):
    lbl(rg,'B%d:C%d'%(r,r),t)
    rg.merge_cells('D%d:G%d'%(r,r))
    style(rg,'D%d:G%d'%(r,r),font=Fo(11,True,AZUL),fill=Fi(CLARO),align=Al('left'),border=BOX)
rg['AL5']='=IFERROR(INDEX($AK$5:$AK$%d,MATCH($D$6,$AJ$5:$AJ$%d,0)),"")'%(4+SHOW,4+SHOW)
rg['D7']='=IF($AL$5="","",INDEX(S_DOC,$AL$5)&"")'
rg['D8']='=IF($AL$5="","",INDEX(S_SEDE,$AL$5)&"   —   "&INDEX(S_PROG,$AL$5))'
rg['D9']='=IF($AL$5="","",INDEX(S_MAIL,$AL$5)&"")'
for k in range(SLOTS):
    r=5+k
    if k==0:
        rg['AE%d'%r]='=IF(TRIM($D$5&"")="","",IFERROR(MATCH("*"&TRIM($D$5)&"*",S_NOM,0),""))'
    else:
        rg['AE%d'%r]=('=IF($AE{p}="","",IFERROR($AE{p}+MATCH("*"&TRIM($D$5)&"*",'
                      'INDEX(S_NOM,$AE{p}+1):INDEX(S_NOM,ROWS(S_NOM)),0),""))').format(p=r-1)
    rg['AF%d'%r]='=IF($AE{r}="","",INDEX(S_DOC,$AE{r})&"")'.format(r=r)
    rg['AG%d'%r]='=IF($AF{r}="","",IF(MATCH($AF{r},$AF$5:$AF${e},0)=ROW()-4,1,0))'.format(r=r,e=4+SLOTS)
    rg['AH%d'%r]='=IF($AG{r}=1,COUNTIF($AG$5:$AG{r},1),"")'.format(r=r)
for k in range(SHOW):
    r=5+k
    rg['AK%d'%r]=('=IFERROR(INDEX($AE$5:$AE${e},MATCH(ROWS($AK$5:AK{r}),$AH$5:$AH${e},0)),"")'
                  ).format(r=r,e=4+SLOTS)
    rg['AJ%d'%r]='=IF($AK{r}="","",INDEX(S_NOM,$AK{r})&"   —   "&INDEX(S_DOC,$AK{r}))'.format(r=r)

# ---- resumen de participación ----
band(rg,'I4:O4','RESUMEN DE PARTICIPACIÓN',AZUL2,'FFFFFF',11)
lbl(rg,'I5:J5','Filtrar por sede:')
rg['K5']='TODAS'
style(rg,'K5:K5',font=Fo(11,True,'7F6000'),fill=Fi(AMAR),align=Al('center'),border=BOX)
dv(rg,'"TODAS,%s"'%','.join(SEDES),['K5'])
rg.merge_cells('L5:O5'); rg['L5']='Cuenta los registros de la tabla de abajo.'
style(rg,'L5:O5',font=Fo(9,False,'808080',True),align=Al('left'))
for ref,t in [('I6:J6','Tipo de participante'),('K6:K6','Participantes'),('L6:L6','%'),
              ('M6:N6','Participaciones'),('O6:O6','%')]:
    if ':' in ref and ref.split(':')[0]!=ref.split(':')[1]: rg.merge_cells(ref)
    rg[ref.split(':')[0]]=t
style(rg,'I6:O6',font=Fo(10,True,'FFFFFF'),fill=Fi('4472C4'),align=Al('center','center',True),border=BOX)
for k,t in enumerate(TIPOS):
    r=7+k
    rg.merge_cells('I%d:J%d'%(r,r)); rg['I%d'%r]=t
    rg['K%d'%r]=('=IF($K$5="TODAS",COUNTIFS(REG_PRIM,1,REG_TIPO,"{t}"),'
                 'COUNTIFS(REG_PRIM,1,REG_TIPO,"{t}",REG_SEDE,$K$5))').format(t=t)
    rg['L%d'%r]='=IFERROR($K{r}/$K$11,"")'.format(r=r)
    rg.merge_cells('M%d:N%d'%(r,r))
    rg['M%d'%r]=('=IF($K$5="TODAS",COUNTIFS(REG_QPART,1,REG_TIPO,"{t}"),'
                 'COUNTIFS(REG_QPART,1,REG_TIPO,"{t}",REG_SEDE,$K$5))').format(t=t)
    rg['O%d'%r]='=IFERROR($M{r}/$M$11,"")'.format(r=r)
rg.merge_cells('I11:J11'); rg['I11']='TOTAL'
rg['K11']='=SUM(K7:K10)'; rg['L11']='=IF($K$11=0,"",SUM(L7:L10))'
rg.merge_cells('M11:N11'); rg['M11']='=SUM(M7:M10)'; rg['O11']='=IF($M$11=0,"",SUM(O7:O10))'
style(rg,'I7:O10',font=Fo(10),fill=Fi('FFFFFF'),align=Al('center'),border=BOX)
style(rg,'I7:J10',align=Al('left'))
style(rg,'I11:O11',font=Fo(10,True),fill=Fi(CLARO),align=Al('center'),border=BOX)
style(rg,'I11:J11',align=Al('left'))
style(rg,'L7:L11',fmt='0.0%'); style(rg,'O7:O11',fmt='0.0%')
style(rg,'K7:K11',fmt='#,##0'); style(rg,'M7:M11',fmt='#,##0')

# ---- tabla de registros ----
HDR={'A':'Q-Part','B':'Fecha','C':'Documento\n(C.C. o ID)','D':'SEDE','E':'APELLIDOS Y NOMBRES',
     'F':'PROGRAMA / ÁREA','G':'CORREO INSTITUCIONAL\n(@uniminuto.edu)','H':'CORREO ADICIONAL',
     'I':'TELÉFONO','J':'TELÉFONO ADICIONAL','K':'TIPO DE\nPARTICIPANTE',
     'L':'Correo','M':'Tipo','N':'Sede','O':'Programa / Área',
     'P':'Actividad / espacio','Q':'Observaciones'}
for c,t in HDR.items(): rg[c+'13']=t
style(rg,'A13:Q13',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center','center',True),border=BOX)
style(rg,'L13:O13',fill=Fi('BF8F00'))
band(rg,'L12:O12','COMPLETE SOLO SI EL NOMBRE APARECE COMO «INEXISTENTE»','BF8F00','FFFFFF',9,'center')
rg.row_dimensions[13].height=36
for i in range(19,31):        # solo S..AD; AE..AL las usa el buscador (filas 5-34)
    rg.cell(row=13,column=i,value='auxiliar - no modificar').font=Fo(8,False,'A6A6A6',True)

for r in range(FIRST,LAST+1):
    rg['A%d'%r]='=IF(OR($S{r}="",$B{r}=""),"",1)'.format(r=r)
    rg['D%d'%r]=('=IF($S{r}="","",UPPER(IF($U{r}<>"",INDEX(BD_EST_SEDE,$U{r})&"",'
                 'IF($T{r}<>"",INDEX(BD_COL_SEDE,$T{r})&"",TRIM($N{r}&"")))))').format(r=r)
    rg['E%d'%r]=('=IF($S{r}="","",UPPER(IF($U{r}<>"",INDEX(BD_EST_NOM,$U{r})&"",'
                 'IF($T{r}<>"",INDEX(BD_COL_NOM,$T{r})&"","INEXISTENTE"))))').format(r=r)
    rg['F%d'%r]=('=IF($S{r}="","",UPPER(IF($U{r}<>"",INDEX(BD_EST_PROG,$U{r})&"",'
                 'IF($T{r}<>"",INDEX(BD_COL_PROG,$T{r})&"",TRIM($O{r}&"")))))').format(r=r)
    rg['G%d'%r]=('=IF($S{r}="","",IF($U{r}<>"",'
                 'IF(INDEX(BD_EST_MAIL,$U{r})&""<>"",INDEX(BD_EST_MAIL,$U{r})&"",TRIM($L{r}&"")),'
                 'IF($T{r}<>"",IF(INDEX(BD_COL_MAIL,$T{r})&""<>"",INDEX(BD_COL_MAIL,$T{r})&"",'
                 'TRIM($L{r}&"")),TRIM($L{r}&""))))').format(r=r)
    rg['H%d'%r]='=IF(OR($X{r}="",$X{r}=$G{r}),"",$X{r})'.format(r=r)
    rg['I%d'%r]=('=IF($S{r}="","",IF($U{r}<>"",INDEX(BD_EST_CEL,$U{r})&"",'
                 'IF($T{r}<>"",INDEX(BD_COL_TEL,$T{r})&"","")))').format(r=r)
    rg['J%d'%r]=('=IF(OR($S{r}="",$U{r}=""),"",'
                 'IF(AND(INDEX(BD_EST_TEL2,$U{r})&""<>"",INDEX(BD_EST_TEL2,$U{r})&""<>$I{r}),'
                 'INDEX(BD_EST_TEL2,$U{r})&"",'
                 'IF(AND(INDEX(BD_EST_TEL3,$U{r})&""<>"",INDEX(BD_EST_TEL3,$U{r})&""<>$I{r}),'
                 'INDEX(BD_EST_TEL3,$U{r})&"","")))').format(r=r)
    rg['K%d'%r]=('=IF($S{r}="","",IF($U{r}<>"",IF(INDEX(BD_EST_COD,$U{r})=1,"ESTUDIANTE",""),'
                 'IF($T{r}<>"",IF(INDEX(BD_COL_COD,$T{r})=27,"PROFESOR",'
                 'IF(INDEX(BD_COL_COD,$T{r})=28,"ADMINISTRATIVO","")),UPPER(TRIM($M{r}&"")))))').format(r=r)
    # auxiliares
    rg['S%d'%r]=('=IF(TRIM($C{r}&"")="","",IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('
                 'TRIM($C{r}&"")," ",""),".",""),",",""),TRIM($C{r}&"")))').format(r=r)
    rg['T%d'%r]=('=IF($S{r}="","",IFERROR(MATCH($S{r},BD_COL_CC,0),'
                 'IFERROR(MATCH($S{r},BD_COL_ID,0),"")))').format(r=r)
    rg['U%d'%r]=('=IF(OR($S{r}="",$T{r}<>""),"",IFERROR(MATCH($S{r},BD_EST_DOC,0),'
                 'IFERROR(MATCH($S{r},BD_EST_ID,0),"")))').format(r=r)
    rg['V%d'%r]='=IF($S{r}="","",$S{r}&"")'.format(r=r)
    rg['W%d'%r]=('=IF($V{r}="","",IF(MATCH($V{r},$V${f}:$V${l},0)=ROW()-{o},1,0))'
                 ).format(r=r,f=FIRST,l=LAST,o=FIRST-1)
    rg['X%d'%r]=('=IF($S{r}="","",IF($U{r}<>"",IFERROR(LEFT(INDEX(BD_EST_MAIL2,$U{r})&"",'
                 'FIND("#",INDEX(BD_EST_MAIL2,$U{r})&"")-1),INDEX(BD_EST_MAIL2,$U{r})&""),'
                 'IF($T{r}<>"",INDEX(BD_COL_MAIL2,$T{r})&"","")))').format(r=r)
    rg['Y%d'%r]='=IF(AND($K{r}="ESTUDIANTE",$F{r}<>""),$F{r}&"","")'.format(r=r)
    rg['Z%d'%r]=('=IF($Y{r}="","",IF(MATCH($Y{r},$Y${f}:$Y${l},0)=ROW()-{o},$Y{r},""))'
                 ).format(r=r,f=FIRST,l=LAST,o=FIRST-1)
    rg['AA%d'%r]='=IF($Z{r}="","",COUNTIF($Z${f}:$Z{r},"?*"))'.format(r=r,f=FIRST)
    rg['AB%d'%r]=('=IF(AND(OR($K{r}="PROFESOR",$K{r}="ADMINISTRATIVO"),$F{r}<>""),$F{r}&"","")'
                  ).format(r=r)
    rg['AC%d'%r]=('=IF($AB{r}="","",IF(MATCH($AB{r},$AB${f}:$AB${l},0)=ROW()-{o},$AB{r},""))'
                  ).format(r=r,f=FIRST,l=LAST,o=FIRST-1)
    rg['AD%d'%r]='=IF($AC{r}="","",COUNTIF($AC${f}:$AC{r},"?*"))'.format(r=r,f=FIRST)

f_auto=Fo(10); f_in=Fo(10,False,'7F6000')
for r in range(FIRST,LAST+1):
    for col in 'ADEFGHIJK':
        c=rg[col+str(r)]; c.font=f_auto; c.fill=Fi(CLARO); c.border=BOX
        c.alignment=Al('left') if col in 'EFGH' else Al('center')
    for col in ('B','C','L','M','N','O','P','Q'):
        c=rg[col+str(r)]; c.font=f_in; c.fill=Fi(AMAR); c.border=BOX
        c.alignment=Al('center') if col in ('B','C','M','N') else Al('left')
    rg['A%d'%r].number_format='0'; rg['B%d'%r].number_format='DD/MM/YYYY'
    rg['C%d'%r].number_format='0'

old=D['reg']; n=0
for i in range(18,min(len(old),2000)):
    o=old[i]
    if o[5] in (None,''): continue
    r=FIRST+n; n+=1
    if r>LAST: break
    rg['B%d'%r]=o[4]; rg['C%d'%r]=o[5]
    if o[16] not in (None,''): rg['P%d'%r]=o[16]
    if o[11] not in (None,''): rg['L%d'%r]=o[11]
    if o[12] not in (None,''): rg['M%d'%r]=str(o[12]).upper()
    if o[13] not in (None,''): rg['N%d'%r]=o[13]
    if o[14] not in (None,''): rg['O%d'%r]=o[14]
print('registros migrados:',n)

dv(rg,'"%s"'%','.join(TIPOS),['M%d:M%d'%(FIRST,LAST)])
dv(rg,'"%s"'%','.join(SEDES),['N%d:N%d'%(FIRST,LAST)])
v=dv(rg,'DATE(2015,1,1)',['B%d:B%d'%(FIRST,LAST)],tipo='date',operator='greaterThan')
v.errorStyle='warning'; v.showErrorMessage=True
rg.conditional_formatting.add('A%d:K%d'%(FIRST,LAST),
    FormulaRule(formula=['$E%d="INEXISTENTE"'%FIRST],fill=Fi('FFD9CC'),stopIfTrue=False))
rg.freeze_panes='D14'
rg.auto_filter.ref='A13:Q%d'%LAST

# ====================== RESULTADOS POR PROGRAMA ======================
rp=wb.create_sheet('Resultados por programa'); rp.sheet_properties.tabColor=AZUL2
rp.sheet_view.showGridLines=False
for k,v in {'A':2,'B':42,'C':16,'D':16,'E':12}.items(): rp.column_dimensions[k].width=v
for i in range(7,16):
    L=CL(i); rp.column_dimensions[L].width=12; rp.column_dimensions[L].hidden=True
band(rp,'B1:E1','PARTICIPACIÓN POR PROGRAMA Y ÁREA',AZUL,'FFFFFF',14)
rp.row_dimensions[1].height=26
lbl(rp,'B3:B3','Sede:')
rp['C3']="=Registros!$K$5"
style(rp,'C3:C3',font=Fo(12,True,AZUL),fill=Fi(CLARO),align=Al('center'),border=BOX)
rp.merge_cells('D3:E3'); rp['D3']='(el filtro se cambia en la hoja Registros)'
style(rp,'D3:E3',font=Fo(9,False,'808080',True),align=Al('left'))

NP=40
def bloque_prog(top, titulo, encab, pn, ix, criterio_tipo, aux):
    """aux = (col programa, col participantes, col clave, col posición)"""
    ap,ac,ak,ao = aux
    band(rp,'B%d:E%d'%(top,top), titulo, AZUL2,'FFFFFF',11)
    for j,t in enumerate([encab,'Participantes','Participaciones','% del total']):
        rp.cell(row=top+1, column=2+j, value=t)
    style(rp,'B%d:E%d'%(top+1,top+1),font=Fo(10,True,'FFFFFF'),fill=Fi('4472C4'),
          align=Al('center','center',True),border=BOX)
    h0, h1 = top+2, top+1+NP
    cnt = lambda crit, ref: ' + '.join(
        ('IF($C$3="TODAS",COUNTIFS({c},1,REG_TIPO,"{t}",REG_PROG,{r}),'
         'COUNTIFS({c},1,REG_TIPO,"{t}",REG_PROG,{r},REG_SEDE,$C$3))').format(c=crit,t=t,r=ref)
        for t in criterio_tipo)
    for k in range(NP):
        r = h0+k
        rp['%s%d'%(ap,r)]=('=IFERROR(INDEX({pn},MATCH(ROWS(${a}${h}:{a}{r}),{ix},0))&"","")'
                           ).format(pn=pn,ix=ix,a=ap,h=h0,r=r)
        rp['%s%d'%(ac,r)]='=IF(${a}{r}="","",{f})'.format(a=ap,r=r,f=cnt('REG_PRIM','$%s%d'%(ap,r)))
        rp['%s%d'%(ak,r)]='=IF(${a}{r}="","",${c}{r}*10000+(10000-ROW()))'.format(a=ap,c=ac,r=r)
        rp['%s%d'%(ao,r)]=('=IFERROR(MATCH(LARGE(${k}${h}:${k}${e},ROWS(${o}${h}:{o}{r})),'
                           '${k}${h}:${k}${e},0),"")').format(k=ak,o=ao,h=h0,e=h1,r=r)
        rp['B%d'%r]='=IF(${o}{r}="","",INDEX(${a}${h}:${a}${e},${o}{r})&"")'.format(o=ao,a=ap,h=h0,e=h1,r=r)
        rp['C%d'%r]='=IF(${o}{r}="","",INDEX(${c}${h}:${c}${e},${o}{r}))'.format(o=ao,c=ac,h=h0,e=h1,r=r)
        rp['D%d'%r]='=IF($B{r}="","",{f})'.format(r=r,f=cnt('REG_QPART','$B%d'%r))
        rp['E%d'%r]='=IFERROR($C{r}/$C${t},"")'.format(r=r,t=h1+1)
    rtot=h1+1
    rp['B%d'%rtot]='TOTAL'
    rp['C%d'%rtot]='=SUM(C%d:C%d)'%(h0,h1); rp['D%d'%rtot]='=SUM(D%d:D%d)'%(h0,h1)
    rp['E%d'%rtot]='=IF($C$%d=0,"",SUM(E%d:E%d))'%(rtot,h0,h1)
    style(rp,'B%d:E%d'%(h0,h1),font=Fo(10),fill=Fi('FFFFFF'),align=Al('center'),border=BOX)
    style(rp,'B%d:B%d'%(h0,h1),align=Al('left'))
    style(rp,'B%d:E%d'%(rtot,rtot),font=Fo(10,True),fill=Fi(CLARO),align=Al('center'),border=BOX)
    style(rp,'B%d:B%d'%(rtot,rtot),align=Al('left'))
    style(rp,'C%d:D%d'%(h0,rtot),fmt='#,##0'); style(rp,'E%d:E%d'%(h0,rtot),fmt='0.0%')
    return rtot

fin1 = bloque_prog(5, 'PROGRAMAS ACADÉMICOS  —  estudiantes', 'Programa académico',
                   'REG_PROGEST','REG_IEST', ['ESTUDIANTE'], ('G','H','I','J'))
bloque_prog(fin1+3, 'ÁREAS Y DEPENDENCIAS  —  profesores y administrativos', 'Área / dependencia',
            'REG_PROGCOL','REG_ICOL', ['PROFESOR','ADMINISTRATIVO'], ('K','L','M','N'))
rp.freeze_panes='A7'
# ===================== EVALUACIÓN =====================
ev = wb.create_sheet('Evaluación'); ev.sheet_properties.tabColor='548235'
ev.sheet_view.showGridLines=False
for k,v in {'A':3,'L':3,'M':9,'N':7,'O':4,'S':4,'Y':4,'AB':13,'AC':3}.items():
    ev.column_dimensions[k].width=v
for c in 'BCDEFGHIJK': ev.column_dimensions[c].width=11
for c in ['P','Q','R','T','U','V','W','X','Z','AA']: ev.column_dimensions[c].width=6
band(ev,'A1:AB1','REGISTRO DE EVALUACIONES', AZUL,'FFFFFF',13)
ev.row_dimensions[1].height=24
ev.merge_cells('O3:AA4')
ev['O3']=('E - Excelente (4)     N - Notable (3)     A - Aceptable (2)     '
          'N/M - Necesita Mejoramiento (1)     N/A - No Aplica (0)')
style(ev,'O3:AA4',font=Fo(10,True,AZUL),fill=Fi('E2EFDA'),align=Al('center','center',True),border=BOX)
ev.merge_cells('M5:N5'); ev['M5']='Promedio'
ev.merge_cells('M6:N6'); ev['M6']='%'
style(ev,'M5:N6',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
ev['AB5']='Promedio %'
style(ev,'AB5:AB5',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center',w=True),border=BOX)

MIRROR={'B':'P','C':'Q','D':'R','E':'T','F':'U','G':'V','H':'W','I':'X','J':'Z','K':'AA'}
for s,d in MIRROR.items():
    ev['%s5'%d]='=IFERROR(AVERAGEIF({d}{f}:{d}{l},">0"),"")'.format(d=d,f=EVA_FIRST,l=EVA_LAST)
    ev['%s6'%d]='=IF({d}5<>"",{d}5*0.25,"")'.format(d=d)
ev['AB6']='=IFERROR(AVERAGE(P6:R6,T6:X6,Z6:AA6),"")'
style(ev,'P5:AA6',font=Fo(9),fill=Fi(CLARO),align=Al('center'),border=BOX)
style(ev,'P5:AA5',fmt='0.00'); style(ev,'P6:AA6',fmt='0%')
style(ev,'AB6:AB6',font=Fo(12,True,AZUL),fill=Fi('E2EFDA'),align=Al('center'),border=BOX,fmt='0.0%')
ev.merge_cells('B6:K6')
ev['B6']='Ingrese aquí los resultados de las evaluaciones (puede copiar y pegar)'
style(ev,'B6:K6',font=Fo(10,True,'7F6000'),fill=Fi(AMAR),align=Al('left'))
for ref,txt in [('B8:D8','I. Planeación'),('E8:I8','II. Desarrollo'),('J8:K8','III. Aportes'),
                ('P8:R8','I. Planeación'),('T8:X8','II. Desarrollo'),('Z8:AA8','III. Aportes')]:
    band(ev,ref,txt,'4472C4','FFFFFF',10,'center')
for j in range(10): ev.cell(row=9, column=2+j, value=j+1)
ev['N9']='Nº'; ev['O9']='I'; ev['S9']='II'; ev['Y9']='III'
for j,d in enumerate(['P','Q','R','T','U','V','W','X','Z','AA']): ev['%s9'%d]=j+1
style(ev,'B9:K9',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
style(ev,'N9:AA9',font=Fo(9,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
CODES='{"0","1","2","3","4","N/A","NA","N/M","NM","A","N","E"}'
for r in range(EVA_FIRST, EVA_LAST+1):
    ev['N%d'%r]=r-EVA_FIRST+1
    for s,d in MIRROR.items():
        ev['%s%d'%(d,r)]=('=IFERROR(IF(TRIM(${s}{r}&"")="","",'
            'IF(ISNUMBER(--TRIM(${s}{r}&"")),'
            'IF(AND(--TRIM(${s}{r}&"")>=0,--TRIM(${s}{r}&"")<=4),--TRIM(${s}{r}&""),"-"),'
            'CHOOSE(MATCH(UPPER(TRIM(${s}{r}&"")),{{"N/A","NA","N/M","NM","A","N","E"}},0),'
            '0,0,1,1,2,3,4))),"-")').format(s=s,r=r)
style(ev,'B%d:K%d'%(EVA_FIRST,EVA_LAST),font=Fo(11),fill=Fi(AMAR),align=Al('center'),border=BOX)
style(ev,'N%d:N%d'%(EVA_FIRST,EVA_LAST),font=Fo(9),fill=Fi(GRIS),align=Al('center'),border=BOX)
style(ev,'P%d:AA%d'%(EVA_FIRST,EVA_LAST),font=Fo(9),fill=Fi('FFFFFF'),align=Al('center'),border=BOX)
for i,row in enumerate(D['eva']):
    if i<9: continue
    r=EVA_FIRST+(i-9)
    if r>EVA_LAST: break
    for j,col in enumerate('BCDEFGHIJK'):
        v2=row[1+j] if 1+j<len(row) else None
        if v2 not in (None,''): ev['%s%d'%(col,r)]=v2
vev=dv(ev,'"E,N,A,N/M,N/A,4,3,2,1,0"',['B%d:K%d'%(EVA_FIRST,EVA_LAST)])
vev.errorStyle='warning'; vev.showErrorMessage=True
vev.promptTitle='Calificación'; vev.showInputMessage=True
vev.prompt='E=4  N=3  A=2  N/M=1  N/A=0. También puede pegar números.'
ev.conditional_formatting.add('B%d:K%d'%(EVA_FIRST,EVA_LAST),
    FormulaRule(formula=['AND(TRIM(B%d&"")<>"",ISERROR(MATCH(UPPER(TRIM(B%d&"")),%s,0)))'
                         %(EVA_FIRST,EVA_FIRST,CODES)],
                fill=Fi('FFC7CE'), font=Font(color='9C0006',bold=True), stopIfTrue=False))
ev.freeze_panes='B10'


# ====================== GRÁFICOS EVALUACIÓN ======================
gr=wb.create_sheet('Gráficos Evaluación'); gr.sheet_properties.tabColor='548235'
gr.sheet_view.showGridLines=False
for k,v in {'A':2,'B':5,'C':34,'D':11}.items(): gr.column_dimensions[k].width=v
band(gr,'B1:D1','RESULTADOS DE LA EVALUACIÓN',AZUL,'FFFFFF',14)
gr.row_dimensions[1].height=26
PREG=[('Divulgación de la actividad','P6','1. La divulgación de la actividad, evento o servicio.'),
 ('Cumplimiento de los tiempos','Q6','2. La actividad, evento o servicio se realizó de acuerdo con los tiempos establecidos.'),
 ('Recursos tecnológicos','R6','3. Utilización de los recursos tecnológicos y audiovisuales.'),
 ('Claridad de la temática','T6','4. La claridad en la temática o propósito de la actividad, evento o servicio.'),
 ('Participación de asistentes','U6','5. La actividad, evento o servicio promovió la participación activa de los asistentes.'),
 ('Actitud del facilitador','V6','6. Actitud y disponibilidad del facilitador o colaborador.'),
 ('Manejo del tema','W6','7. El manejo del tema por parte del facilitador o colaborador.'),
 ('Lugar y condiciones físicas','X6','8. El lugar y sus condiciones físicas.'),
 ('Cumplió las expectativas','Z6','9. Se cumplió su expectativa frente a la actividad, evento o servicio.'),
 ('Aporte a la formación','AA6','10. La actividad, evento o servicio realizado aportó a su formación integral.')]
gr.merge_cells('B3:C3'); gr['B3']='Aspecto evaluado'; gr['D3']='%'
style(gr,'B3:D3',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
for k,(corta,src_,larga) in enumerate(PREG):
    r=4+k
    gr['B%d'%r]=k+1; gr['C%d'%r]=corta; gr['D%d'%r]='=Evaluación!%s'%src_
style(gr,'B4:D13',font=Fo(10),fill=Fi('FFFFFF'),align=Al('left'),border=BOX)
style(gr,'B4:B13',align=Al('center'))
style(gr,'D4:D13',font=Fo(10,True,AZUL),fill=Fi(CLARO),align=Al('center'),fmt='0.0%')
gr.merge_cells('B15:C15'); gr['B15']='PROMEDIO TOTAL'; gr['D15']='=Evaluación!AB6'
style(gr,'B15:D15',font=Fo(12,True,'FFFFFF'),fill=Fi(AZUL),align=Al('center'),border=BOX)
style(gr,'D15:D15',fmt='0.0%')

ch=BarChart(); ch.type='bar'; ch.style=None
ch.title=None; ch.legend=None
ch.x_axis.numFmt='0%'; ch.x_axis.scaling.min=0; ch.x_axis.scaling.max=1
ch.x_axis.majorGridlines=None; ch.y_axis.majorGridlines=None
ch.x_axis.delete=False; ch.y_axis.delete=False
ch.gapWidth=45
ch.add_data(Reference(gr,min_col=4,min_row=4,max_row=13), titles_from_data=False)
ch.set_categories(Reference(gr,min_col=3,min_row=4,max_row=13))
ch.series[0].graphicalProperties = GraphicalProperties(solidFill='2E5B9A')
ch.dataLabels=DataLabelList(); ch.dataLabels.showVal=True; ch.dataLabels.numFmt='0.0%'
ch.width=19; ch.height=11
gr.add_chart(ch,'F3')

gr['B18']='Preguntas completas de la encuesta'
gr['B18'].font=Fo(9,True,'808080')
for k,(corta,src_,larga) in enumerate(PREG):
    r=19+k
    gr.merge_cells('B%d:D%d'%(r,r)); gr['B%d'%r]=larga
    style(gr,'B%d:D%d'%(r,r),font=Fo(8,False,'808080'),align=Al('left'))
# ===================== BASES =====================
def base(nombre, filas):
    ws=wb.create_sheet(nombre); ws.sheet_properties.tabColor='808080'
    for f in filas: ws.append(f)
    head=filas[0]; ncol=len(head)
    style(ws,'A1:%s1'%CL(ncol),font=Fo(10,True,'FFFFFF'),fill=Fi('595959'),
          align=Al('center','center',True),border=BOX)
    ws.row_dimensions[1].height=30; ws.freeze_panes='A2'
    ws.auto_filter.ref='A1:%s%d'%(CL(ncol),len(filas))
    for j,h in enumerate(head):
        L=CL(j+1); hu=str(h).upper() if h else ''
        ws.column_dimensions[L].width = 30 if ('NOMBRE' in hu or 'CORREO' in hu or 'PROGRAMA' in hu) else 14
        if 'FECHA' in hu:
            for r in range(2,len(filas)+1): ws.cell(row=r,column=j+1).number_format='DD/MM/YYYY'
base('BD ADM-DOC', D['adm'])
base('BD EST',     D['est'])


# ====================== NOMBRES DEFINIDOS ======================
EST=lambda c:"'BD EST'!${c}$2:${c}${n}".format(c=c,n=EST_LAST)
COL=lambda c:"'BD ADM-DOC'!${c}$2:${c}${n}".format(c=c,n=COL_LAST)
RG =lambda c:"Registros!${c}${f}:${c}${l}".format(c=c,f=FIRST,l=LAST)
NAMES={
 'BD_EST_ID':EST('B'),'BD_EST_NOM':EST('C'),'BD_EST_SEDE':EST('G'),'BD_EST_PROG':EST('K'),
 'BD_EST_CEL':EST('V'),'BD_EST_TEL2':EST('W'),'BD_EST_TEL3':EST('X'),
 'BD_EST_MAIL':EST('Y'),'BD_EST_MAIL2':EST('Z'),'BD_EST_DOC':EST('AB'),'BD_EST_COD':EST('AU'),
 'BD_COL_CC':COL('B'),'BD_COL_ID':COL('C'),'BD_COL_NOM':COL('D'),'BD_COL_COD':COL('I'),
 'BD_COL_PROG':COL('K'),'BD_COL_MAIL':COL('L'),'BD_COL_MAIL2':COL('M'),'BD_COL_TEL':COL('N'),
 'BD_COL_SEDE':COL('P'),
 'REG_QPART':RG('A'),'REG_SEDE':RG('D'),'REG_PROG':RG('F'),'REG_TIPO':RG('K'),
 'REG_PRIM':RG('W'),
 'REG_PROGEST':RG('Z'),'REG_IEST':RG('AA'),
 'REG_PROGCOL':RG('AC'),'REG_ICOL':RG('AD'),
 'S_NOM':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_NOM,BD_COL_NOM)',
 'S_DOC':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_DOC,BD_COL_CC)',
 'S_SEDE':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_SEDE,BD_COL_SEDE)',
 'S_PROG':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_PROG,BD_COL_PROG)',
 'S_MAIL':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_MAIL,BD_COL_MAIL)'}
for n2,v2 in NAMES.items(): wb.defined_names.add(DefinedName(n2,attr_text=v2))
wb.calculation.fullCalcOnLoad=True
wb.properties.title='Registro de Asistencia'
out=SP+'/Registro_Asistencia_2026-1_V12.xlsx'
wb.save(out); print('guardado:',out)
