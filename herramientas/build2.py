# -*- coding: utf-8 -*-
"""Registro de Asistencia V11 — versión simplificada."""
import pickle
SP='/tmp/claude-0/-home-user-uniminuto/b364fe32-f339-5501-a525-09b94e80e92f/scratchpad'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

D = pickle.load(open(SP+'/data.pkl','rb'))

FIRST, NROWS = 6, 2000
LAST = FIRST + NROWS - 1            # 2005
EST_LAST, COL_LAST = 60001, 6001
EVA_FIRST, EVA_LAST = 10, 509
SEDES = ['BUG','BVA','CHI','CIN','FLO','IPI','PAS','PER','QUI']

AZUL='1F3864'; AZUL2='2E5B9A'; CLARO='EAF1FA'; AMAR='FFF2CC'; GRIS='F7F7F7'
thin = Side(style='thin', color='BFBFBF')
BOX  = Border(left=thin,right=thin,top=thin,bottom=thin)
def Fo(sz=10,b=False,c='000000',i=False): return Font(name='Calibri',size=sz,bold=b,color=c,italic=i)
def Fi(c): return PatternFill('solid', fgColor=c)
def Al(h='center',v='center',w=False): return Alignment(horizontal=h,vertical=v,wrap_text=w)

def style(ws, ref, font=None, fill=None, align=None, border=None, fmt=None):
    for row in ws[ref]:
        for c in row:
            if font: c.font=font
            if fill: c.fill=fill
            if align: c.alignment=align
            if border: c.border=border
            if fmt: c.number_format=fmt

def band(ws, ref, text, bg=AZUL, fg='FFFFFF', size=11, h='left'):
    ws.merge_cells(ref); ws[ref.split(':')[0]] = text
    style(ws, ref, font=Fo(size,True,fg), fill=Fi(bg), align=Al(h,'center'))

def dv(ws, f1, cells, tipo='list', **kw):
    v = DataValidation(type=tipo, formula1=f1, allow_blank=True, **kw)
    ws.add_data_validation(v)
    for c in cells: v.add(c)
    return v

wb = Workbook(); wb.remove(wb.active)

# ===================== REGISTROS =====================
rg = wb.create_sheet('Registros'); rg.sheet_properties.tabColor = AZUL
rg.sheet_view.showGridLines = False
for k,v in {'A':7,'B':12,'C':16,'D':8,'E':36,'F':28,'G':32,'H':14,'I':15,
            'J':30,'K':24,'L':24,'M':2}.items():
    rg.column_dimensions[k].width = v
for i in range(14, 31):                      # N..AD auxiliares
    L=CL(i); rg.column_dimensions[L].width=11; rg.column_dimensions[L].hidden=True

band(rg,'A1:L1','REGISTRO DE ASISTENCIA', AZUL, 'FFFFFF', 14)
rg.row_dimensions[1].height = 26
band(rg,'A2:L2','Escriba la cédula o el ID en la columna «Documento». Los demás datos aparecen solos.',
     AZUL2, 'FFFFFF', 10)

# --- buscador de una sola línea ---
rg.merge_cells('B3:C3'); rg['B3']='¿No sabe el documento? Escriba el nombre:'
style(rg,'B3:C3',font=Fo(10,True,AZUL),align=Al('right'))
rg.merge_cells('D3:E3')
style(rg,'D3:E3',font=Fo(11,True,'7F6000'),fill=Fi(AMAR),align=Al('left'),border=BOX)
rg['AC3']='=IF(TRIM($D$3&"")="","",IFERROR(MATCH("*"&TRIM($D$3)&"*",BD_EST_NOM,0),""))'
rg['AD3']=('=IF(OR(TRIM($D$3&"")="",$AC$3<>""),"",'
           'IFERROR(MATCH("*"&TRIM($D$3)&"*",BD_COL_NOM,0),""))')
rg['F3']=('=IF(TRIM($D$3&"")="","",IF($AC$3<>"",INDEX(BD_EST_DOC,$AC$3)&"",'
          'IF($AD$3<>"",INDEX(BD_COL_CC,$AD$3)&"","no encontrado")))')
rg.merge_cells('G3:I3')
rg['G3']=('=IF(TRIM($D$3&"")="","",IF($AC$3<>"",INDEX(BD_EST_NOM,$AC$3)&"",'
          'IF($AD$3<>"",INDEX(BD_COL_NOM,$AD$3)&"","")))')
style(rg,'F3:I3',font=Fo(11,True,AZUL),fill=Fi(CLARO),align=Al('left'),border=BOX)
style(rg,'F3:F3',align=Al('center'))
rg.row_dimensions[3].height = 20

# --- tabla ---
HDR = {'A':'Q-Part','B':'Fecha','C':'Documento\n(C.C. o ID)','D':'Sede','E':'Apellidos y Nombres',
       'F':'Programa / Área','G':'Correo electrónico','H':'Teléfono','I':'Tipo de\nparticipante',
       'J':'Correo\n(escríbalo si dice «Inexistente»)','K':'Actividad / espacio','L':'Observaciones'}
for c,t in HDR.items(): rg[c+'5']=t
style(rg,'A5:L5',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center','center',True),border=BOX)
style(rg,'J5:J5',fill=Fi('BF8F00'))
rg.row_dimensions[5].height = 34
for i in range(14,31):
    rg.cell(row=5, column=i, value='auxiliar - no modificar').font=Fo(8,False,'A6A6A6',True)

BLOQ = {'Estudiante':('S','T','U'),'Docente':('V','W','X'),'Administrativo':('Y','Z','AA')}
for r in range(FIRST, LAST+1):
    rg['A%d'%r]='=IF(OR($N{r}="",$B{r}=""),"",1)'.format(r=r)
    rg['D%d'%r]=('=IF($N{r}="","",IF($P{r}<>"",INDEX(BD_EST_SEDE,$P{r})&"",'
                 'IF($O{r}<>"",INDEX(BD_COL_SEDE,$O{r})&"","")))').format(r=r)
    rg['E%d'%r]=('=IF($N{r}="","",IF($P{r}<>"",INDEX(BD_EST_NOM,$P{r})&"",'
                 'IF($O{r}<>"",INDEX(BD_COL_NOM,$O{r})&"","Inexistente")))').format(r=r)
    rg['F%d'%r]=('=IF($N{r}="","",IF($P{r}<>"",INDEX(BD_EST_PROG,$P{r})&"",'
                 'IF($O{r}<>"",INDEX(BD_COL_PROG,$O{r})&"","")))').format(r=r)
    rg['G%d'%r]=('=IF($N{r}="","",IF(TRIM($J{r}&"")<>"",TRIM($J{r}&""),'
                 'IF($P{r}<>"",INDEX(BD_EST_MAIL,$P{r})&"",'
                 'IF($O{r}<>"",IF(INDEX(BD_COL_MAIL,$O{r})&""<>"",INDEX(BD_COL_MAIL,$O{r})&"",'
                 'INDEX(BD_COL_MAIL2,$O{r})&""),""))))').format(r=r)
    rg['H%d'%r]=('=IF($N{r}="","",IF($P{r}<>"",'
                 'IF(INDEX(BD_EST_CEL,$P{r})&""<>"",INDEX(BD_EST_CEL,$P{r})&"",'
                 'IF(INDEX(BD_EST_TEL2,$P{r})&""<>"",INDEX(BD_EST_TEL2,$P{r})&"",'
                 'INDEX(BD_EST_TEL3,$P{r})&"")),'
                 'IF($O{r}<>"",INDEX(BD_COL_TEL,$O{r})&"","")))').format(r=r)
    rg['I%d'%r]=('=IF($N{r}="","",IF($P{r}<>"","Estudiante",'
                 'IF($O{r}<>"",IF(INDEX(BD_COL_TIPO,$O{r})&""="DOC","Docente","Administrativo"),"")))'
                 ).format(r=r)
    rg['N%d'%r]=('=IF(TRIM($C{r}&"")="","",IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('
                 'TRIM($C{r}&"")," ",""),".",""),",",""),TRIM($C{r}&"")))').format(r=r)
    rg['O%d'%r]=('=IF($N{r}="","",IFERROR(MATCH($N{r},BD_COL_CC,0),'
                 'IFERROR(MATCH($N{r},BD_COL_ID,0),"")))').format(r=r)
    rg['P%d'%r]=('=IF(OR($N{r}="",$O{r}<>""),"",IFERROR(MATCH($N{r},BD_EST_DOC,0),'
                 'IFERROR(MATCH($N{r},BD_EST_ID,0),"")))').format(r=r)
    rg['Q%d'%r]='=IF($N{r}="","",$N{r}&"")'.format(r=r)
    rg['R%d'%r]=('=IF($Q{r}="","",IF(MATCH($Q{r},$Q${f}:$Q${l},0)=ROW()-{o},1,0))'
                 ).format(r=r,f=FIRST,l=LAST,o=FIRST-1)
    for tipo,(kc,pc,ic) in BLOQ.items():
        rg['%s%d'%(kc,r)]='=IF(AND($I{r}="{t}",$F{r}<>""),$F{r}&"","")'.format(r=r,t=tipo)
        rg['%s%d'%(pc,r)]=('=IF(${k}{r}="","",IF(MATCH(${k}{r},${k}${f}:${k}${l},0)=ROW()-{o},${k}{r},""))'
                           ).format(r=r,k=kc,f=FIRST,l=LAST,o=FIRST-1)
        rg['%s%d'%(ic,r)]='=IF(${p}{r}="","",COUNTIF(${p}${f}:${p}{r},"?*"))'.format(r=r,p=pc,f=FIRST)

f_auto=Fo(10); f_in=Fo(10,False,'7F6000')
for r in range(FIRST, LAST+1):
    for col in 'ADEFGHI':
        c=rg[col+str(r)]; c.font=f_auto; c.fill=Fi(CLARO); c.border=BOX
        c.alignment = Al('left') if col in 'EFG' else Al('center')
    for col in ('B','C','J','K','L'):
        c=rg[col+str(r)]; c.font=f_in; c.fill=Fi(AMAR); c.border=BOX
        c.alignment = Al('center') if col in 'BC' else Al('left')
    rg['A%d'%r].number_format='0'
    rg['B%d'%r].number_format='DD/MM/YYYY'
    rg['C%d'%r].number_format='0'

old=D['reg']; n=0
for i in range(18, min(len(old),2000)):
    o=old[i]
    if o[5] in (None,''): continue
    r=FIRST+n; n+=1
    if r>LAST: break
    rg['B%d'%r]=o[4]; rg['C%d'%r]=o[5]
    if o[16] not in (None,''): rg['K%d'%r]=o[16]
    if o[11] not in (None,''): rg['J%d'%r]=o[11]
print('registros migrados:', n)

v=dv(rg,'DATE(2015,1,1)',['B%d:B%d'%(FIRST,LAST)],tipo='date',operator='greaterThan')
v.errorStyle='warning'; v.showErrorMessage=True

rg.conditional_formatting.add('A%d:L%d'%(FIRST,LAST),
    FormulaRule(formula=['$E%d="Inexistente"'%FIRST], fill=Fi('FFD9CC'), stopIfTrue=False))
rg.freeze_panes='D6'
rg.auto_filter.ref='A5:L%d'%LAST

# ===================== RESULTADOS POR TIPO =====================
rt = wb.create_sheet('Resultados por tipo'); rt.sheet_properties.tabColor=AZUL2
rt.sheet_view.showGridLines=False
for k,v in {'A':2,'B':24,'C':15,'D':11}.items(): rt.column_dimensions[k].width=v
band(rt,'B1:D1','RESULTADOS POR TIPO DE PARTICIPANTE', AZUL, 'FFFFFF', 13)
rt.row_dimensions[1].height=24
rt['B3']='Seleccione Sede:'; rt['B3'].font=Fo(11,True,AZUL); rt['B3'].alignment=Al('right')
rt['C3']='TODAS'
style(rt,'C3:C3',font=Fo(12,True,'7F6000'),fill=Fi(AMAR),align=Al('center'),border=BOX)
dv(rt,'"TODAS,%s"'%','.join(SEDES),['C3'])

def tabla(top, titulo, etiqueta, criterio):
    band(rt,'B%d:D%d'%(top,top), titulo, AZUL2, 'FFFFFF', 11)
    for j,t in enumerate(['Tipo', etiqueta, '%']): rt.cell(row=top+1, column=2+j, value=t)
    style(rt,'B%d:D%d'%(top+1,top+1),font=Fo(10,True,'FFFFFF'),fill=Fi('4472C4'),
          align=Al('center'),border=BOX)
    for k,(lab,val) in enumerate([('Estudiantes','Estudiante'),('Administrativos','Administrativo'),
                                  ('Docentes','Docente')]):
        r=top+2+k
        rt['B%d'%r]=lab
        rt['C%d'%r]=('=IF($C$3="TODAS",COUNTIFS({c},1,REG_TIPO,"{v}"),'
                     'COUNTIFS({c},1,REG_TIPO,"{v}",REG_SEDE,$C$3))').format(c=criterio,v=val)
        rt['D%d'%r]='=IFERROR($C{r}/$C${t},"")'.format(r=r,t=top+5)
    rtot=top+5
    rt['B%d'%rtot]='Total'
    rt['C%d'%rtot]='=SUM(C%d:C%d)'%(top+2,top+4)
    rt['D%d'%rtot]='=IF($C$%d=0,"",SUM(D%d:D%d))'%(rtot,top+2,top+4)
    style(rt,'B%d:D%d'%(top+2,top+4),font=Fo(11),fill=Fi('FFFFFF'),align=Al('center'),border=BOX)
    style(rt,'B%d:B%d'%(top+2,top+4),align=Al('left'))
    style(rt,'B%d:D%d'%(rtot,rtot),font=Fo(11,True),fill=Fi(CLARO),align=Al('center'),border=BOX)
    style(rt,'B%d:B%d'%(rtot,rtot),align=Al('left'))
    style(rt,'D%d:D%d'%(top+2,rtot),fmt='0.0%')
    style(rt,'C%d:C%d'%(top+2,rtot),fmt='#,##0')

tabla(5,  'PARTICIPANTES  (personas distintas)',       'Participantes',  'REG_PRIMERA')
tabla(13, 'PARTICIPACIONES  (asistencias registradas)','Participaciones','REG_QPART')

# ===================== RESULTADOS POR PROGRAMA =====================
rp = wb.create_sheet('Resultados por programa'); rp.sheet_properties.tabColor=AZUL2
rp.sheet_view.showGridLines=False
for k,v in {'A':2,'B':30,'C':14,'D':15,'E':3,'F':28,'G':14,'H':15,'I':3,'J':28,'K':14,'L':15}.items():
    rp.column_dimensions[k].width=v
band(rp,'B1:L1','RESULTADOS POR PROGRAMA', AZUL, 'FFFFFF', 13)
rp.row_dimensions[1].height=24
rp['B3']='Sede seleccionada:'; rp['B3'].font=Fo(11,True,AZUL); rp['B3'].alignment=Al('right')
rp['C3']="='Resultados por tipo'!$C$3"
style(rp,'C3:C3',font=Fo(12,True,AZUL),fill=Fi(CLARO),align=Al('center'),border=BOX)
rp['D3']='(se cambia en Resultados por tipo)'; rp['D3'].font=Fo(9,False,'808080',True)
rp['D3'].alignment=Al('left')

TOP, NP = 6, 50
for cini,titulo,tipo,pn,ix in [(2,'ESTUDIANTES','Estudiante','REG_PROGEST','REG_IEST'),
                               (6,'DOCENTES','Docente','REG_PROGDOC','REG_IDOC'),
                               (10,'ADMINISTRATIVOS','Administrativo','REG_PROGADM','REG_IADM')]:
    c0,c1,c2 = CL(cini),CL(cini+1),CL(cini+2)
    band(rp,'%s5:%s5'%(c0,c2), titulo, AZUL2,'FFFFFF',11,'center')
    for j,t in enumerate(['Programa / Área','Participantes','Participaciones']):
        rp.cell(row=TOP, column=cini+j, value=t)
    style(rp,'%s%d:%s%d'%(c0,TOP,c2,TOP),font=Fo(10,True,'FFFFFF'),fill=Fi('4472C4'),
          align=Al('center','center',True),border=BOX)
    for k in range(NP):
        r=TOP+1+k
        rp['%s%d'%(c0,r)]=('=IFERROR(INDEX({pn},MATCH(ROWS(${c}${t}:{c}{r}),{ix},0))&"","")'
                           ).format(pn=pn,ix=ix,c=c0,t=TOP+1,r=r)
        for cc,cr in ((c1,'REG_PRIMERA'),(c2,'REG_QPART')):
            rp['%s%d'%(cc,r)]=('=IF(${c}{r}="","",IF($C$3="TODAS",'
                               'COUNTIFS({cr},1,REG_TIPO,"{tp}",REG_PROG,${c}{r}),'
                               'COUNTIFS({cr},1,REG_TIPO,"{tp}",REG_PROG,${c}{r},REG_SEDE,$C$3)))'
                               ).format(c=c0,r=r,cr=cr,tp=tipo)
    fin=TOP+NP
    style(rp,'%s%d:%s%d'%(c0,TOP+1,c2,fin),font=Fo(10),align=Al('center'),border=BOX)
    style(rp,'%s%d:%s%d'%(c0,TOP+1,c0,fin),align=Al('left'))
    style(rp,'%s%d:%s%d'%(c1,TOP+1,c2,fin),fmt='#,##0')
rp.freeze_panes='A%d'%(TOP+1)

# ===================== EVALUACIÓN =====================
ev = wb.create_sheet('Evaluación'); ev.sheet_properties.tabColor='548235'
ev.sheet_view.showGridLines=False
for k,v in {'A':3,'L':3,'M':9,'N':7,'O':4,'S':4,'Y':4,'AB':13,'AC':3}.items():
    ev.column_dimensions[k].width=v
for c in 'BCDEFGHIJK': ev.column_dimensions[c].width=11
for c in ['P','Q','R','T','U','V','W','X','Z','AA']: ev.column_dimensions[c].width=6
band(ev,'A1:AB1','REGISTRO DE EVALUACIONES', AZUL,'FFFFFF',13)
ev.row_dimensions[1].height=24
ev.merge_cells('O3:AA4')
ev['O3']=('E - Excelente (4)     N - Notable (3)     A - Aceptable (2)     '
          'N/M - Necesita Mejoramiento (1)     N/A - No Aplica (0)')
style(ev,'O3:AA4',font=Fo(10,True,AZUL),fill=Fi('E2EFDA'),align=Al('center','center',True),border=BOX)
ev.merge_cells('M5:N5'); ev['M5']='Promedio'
ev.merge_cells('M6:N6'); ev['M6']='%'
style(ev,'M5:N6',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
ev['AB5']='Promedio %'
style(ev,'AB5:AB5',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center',w=True),border=BOX)

MIRROR={'B':'P','C':'Q','D':'R','E':'T','F':'U','G':'V','H':'W','I':'X','J':'Z','K':'AA'}
for s,d in MIRROR.items():
    ev['%s5'%d]='=IFERROR(AVERAGEIF({d}{f}:{d}{l},">0"),"")'.format(d=d,f=EVA_FIRST,l=EVA_LAST)
    ev['%s6'%d]='=IF({d}5<>"",{d}5*0.25,"")'.format(d=d)
ev['AB6']='=IFERROR(AVERAGE(P6:R6,T6:X6,Z6:AA6),"")'
style(ev,'P5:AA6',font=Fo(9),fill=Fi(CLARO),align=Al('center'),border=BOX)
style(ev,'P5:AA5',fmt='0.00'); style(ev,'P6:AA6',fmt='0%')
style(ev,'AB6:AB6',font=Fo(12,True,AZUL),fill=Fi('E2EFDA'),align=Al('center'),border=BOX,fmt='0.0%')
ev.merge_cells('B6:K6')
ev['B6']='Ingrese aquí los resultados de las evaluaciones (puede copiar y pegar)'
style(ev,'B6:K6',font=Fo(10,True,'7F6000'),fill=Fi(AMAR),align=Al('left'))
for ref,txt in [('B8:D8','I. Planeación'),('E8:I8','II. Desarrollo'),('J8:K8','III. Aportes'),
                ('P8:R8','I. Planeación'),('T8:X8','II. Desarrollo'),('Z8:AA8','III. Aportes')]:
    band(ev,ref,txt,'4472C4','FFFFFF',10,'center')
for j in range(10): ev.cell(row=9, column=2+j, value=j+1)
ev['N9']='Nº'; ev['O9']='I'; ev['S9']='II'; ev['Y9']='III'
for j,d in enumerate(['P','Q','R','T','U','V','W','X','Z','AA']): ev['%s9'%d]=j+1
style(ev,'B9:K9',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
style(ev,'N9:AA9',font=Fo(9,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
CODES='{"0","1","2","3","4","N/A","NA","N/M","NM","A","N","E"}'
for r in range(EVA_FIRST, EVA_LAST+1):
    ev['N%d'%r]=r-EVA_FIRST+1
    for s,d in MIRROR.items():
        ev['%s%d'%(d,r)]=('=IFERROR(IF(TRIM(${s}{r}&"")="","",'
            'IF(ISNUMBER(--TRIM(${s}{r}&"")),'
            'IF(AND(--TRIM(${s}{r}&"")>=0,--TRIM(${s}{r}&"")<=4),--TRIM(${s}{r}&""),"-"),'
            'CHOOSE(MATCH(UPPER(TRIM(${s}{r}&"")),{{"N/A","NA","N/M","NM","A","N","E"}},0),'
            '0,0,1,1,2,3,4))),"-")').format(s=s,r=r)
style(ev,'B%d:K%d'%(EVA_FIRST,EVA_LAST),font=Fo(11),fill=Fi(AMAR),align=Al('center'),border=BOX)
style(ev,'N%d:N%d'%(EVA_FIRST,EVA_LAST),font=Fo(9),fill=Fi(GRIS),align=Al('center'),border=BOX)
style(ev,'P%d:AA%d'%(EVA_FIRST,EVA_LAST),font=Fo(9),fill=Fi('FFFFFF'),align=Al('center'),border=BOX)
for i,row in enumerate(D['eva']):
    if i<9: continue
    r=EVA_FIRST+(i-9)
    if r>EVA_LAST: break
    for j,col in enumerate('BCDEFGHIJK'):
        v2=row[1+j] if 1+j<len(row) else None
        if v2 not in (None,''): ev['%s%d'%(col,r)]=v2
vev=dv(ev,'"E,N,A,N/M,N/A,4,3,2,1,0"',['B%d:K%d'%(EVA_FIRST,EVA_LAST)])
vev.errorStyle='warning'; vev.showErrorMessage=True
vev.promptTitle='Calificación'; vev.showInputMessage=True
vev.prompt='E=4  N=3  A=2  N/M=1  N/A=0. También puede pegar números.'
ev.conditional_formatting.add('B%d:K%d'%(EVA_FIRST,EVA_LAST),
    FormulaRule(formula=['AND(TRIM(B%d&"")<>"",ISERROR(MATCH(UPPER(TRIM(B%d&"")),%s,0)))'
                         %(EVA_FIRST,EVA_FIRST,CODES)],
                fill=Fi('FFC7CE'), font=Font(color='9C0006',bold=True), stopIfTrue=False))
ev.freeze_panes='B10'

# ===================== GRÁFICOS EVALUACIÓN =====================
gr = wb.create_sheet('Gráficos Evaluación'); gr.sheet_properties.tabColor='548235'
gr.sheet_view.showGridLines=False
for k,v in {'A':2,'B':5,'C':30,'D':10,'E':78}.items(): gr.column_dimensions[k].width=v
band(gr,'B1:E1','RESULTADOS DE LA EVALUACIÓN', AZUL,'FFFFFF',13)
gr.row_dimensions[1].height=24

PREG = [
 ('Divulgación de la actividad','P6','I. Planeación',
  '1. La divulgación de la actividad, evento o servicio.'),
 ('Cumplimiento de los tiempos','Q6','I. Planeación',
  '2. La actividad, evento o servicio se realizó de acuerdo con los tiempos establecidos.'),
 ('Recursos tecnológicos','R6','I. Planeación',
  '3. Utilización de los recursos tecnológicos y audiovisuales.'),
 ('Claridad de la temática','T6','II. Desarrollo',
  '4. La claridad en la temática o propósito de la actividad, evento o servicio.'),
 ('Participación de asistentes','U6','II. Desarrollo',
  '5. La actividad, evento o servicio promovió la participación activa de los asistentes.'),
 ('Actitud del facilitador','V6','II. Desarrollo',
  '6. Actitud y disponibilidad del facilitador o colaborador.'),
 ('Manejo del tema','W6','II. Desarrollo',
  '7. El manejo del tema por parte del facilitador o colaborador.'),
 ('Lugar y condiciones físicas','X6','II. Desarrollo',
  '8. El lugar y sus condiciones físicas.'),
 ('Cumplió las expectativas','Z6','III. Aportes',
  '9. Se cumplió su expectativa frente a la actividad, evento o servicio.'),
 ('Aporte a la formación','AA6','III. Aportes',
  '10. La actividad, evento o servicio realizado aportó a su formación integral.'),
]
for j,t in enumerate(['Nº','Aspecto evaluado','%','Pregunta completa']):
    gr.cell(row=3, column=2+j, value=t)
style(gr,'B3:E3',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
for k,(corta,src,sec,larga) in enumerate(PREG):
    r=4+k
    gr['B%d'%r]=k+1
    gr['C%d'%r]=corta
    gr['D%d'%r]='=Evaluación!%s'%src
    gr['E%d'%r]=larga
style(gr,'B4:E13',font=Fo(10),fill=Fi('FFFFFF'),align=Al('left'),border=BOX)
style(gr,'B4:B13',align=Al('center'))
style(gr,'D4:D13',font=Fo(11,True,AZUL),fill=Fi(CLARO),align=Al('center'),fmt='0.0%')
gr.merge_cells('B15:C15'); gr['B15']='PROMEDIO TOTAL'
gr['D15']='=Evaluación!AB6'
style(gr,'B15:D15',font=Fo(12,True,'FFFFFF'),fill=Fi(AZUL),align=Al('center'),border=BOX)
style(gr,'D15:D15',fmt='0.0%')

ch = BarChart(); ch.type='bar'; ch.style=10
ch.title='Resultado por aspecto evaluado'
ch.x_axis.numFmt='0%'
ch.x_axis.scaling.min=0; ch.x_axis.scaling.max=1
ch.y_axis.delete=False; ch.x_axis.delete=False
ch.legend=None
ch.add_data(Reference(gr,min_col=4,min_row=4,max_row=13), titles_from_data=False)
ch.set_categories(Reference(gr,min_col=3,min_row=4,max_row=13))
ch.dataLabels=DataLabelList(); ch.dataLabels.showVal=True
ch.width=24; ch.height=13
gr.add_chart(ch,'B17')

# ===================== BASES =====================
def base(nombre, filas):
    ws=wb.create_sheet(nombre); ws.sheet_properties.tabColor='808080'
    for f in filas: ws.append(f)
    head=filas[0]; ncol=len(head)
    style(ws,'A1:%s1'%CL(ncol),font=Fo(10,True,'FFFFFF'),fill=Fi('595959'),
          align=Al('center','center',True),border=BOX)
    ws.row_dimensions[1].height=30; ws.freeze_panes='A2'
    ws.auto_filter.ref='A1:%s%d'%(CL(ncol),len(filas))
    for j,h in enumerate(head):
        L=CL(j+1); hu=str(h).upper() if h else ''
        ws.column_dimensions[L].width = 30 if ('NOMBRE' in hu or 'CORREO' in hu or 'PROGRAMA' in hu) else 14
        if 'FECHA' in hu:
            for r in range(2,len(filas)+1): ws.cell(row=r,column=j+1).number_format='DD/MM/YYYY'
base('BD ADM-DOC', D['adm'])
base('BD EST',     D['est'])

# ===================== NOMBRES =====================
EST=lambda c:"'BD EST'!${c}$2:${c}${n}".format(c=c,n=EST_LAST)
COL=lambda c:"'BD ADM-DOC'!${c}$2:${c}${n}".format(c=c,n=COL_LAST)
RG =lambda c:"Registros!${c}${f}:${c}${l}".format(c=c,f=FIRST,l=LAST)
NAMES={'BD_EST_ID':EST('B'),'BD_EST_NOM':EST('C'),'BD_EST_SEDE':EST('G'),'BD_EST_PROG':EST('K'),
 'BD_EST_CEL':EST('V'),'BD_EST_TEL2':EST('W'),'BD_EST_TEL3':EST('X'),'BD_EST_MAIL':EST('Y'),
 'BD_EST_DOC':EST('AB'),
 'BD_COL_CC':COL('B'),'BD_COL_ID':COL('C'),'BD_COL_NOM':COL('D'),'BD_COL_TIPO':COL('J'),
 'BD_COL_PROG':COL('K'),'BD_COL_MAIL':COL('L'),'BD_COL_MAIL2':COL('M'),'BD_COL_TEL':COL('N'),
 'BD_COL_SEDE':COL('P'),
 'REG_QPART':RG('A'),'REG_SEDE':RG('D'),'REG_PROG':RG('F'),'REG_TIPO':RG('I'),
 'REG_PRIMERA':RG('R'),
 'REG_PROGEST':RG('T'),'REG_IEST':RG('U'),
 'REG_PROGDOC':RG('W'),'REG_IDOC':RG('X'),
 'REG_PROGADM':RG('Z'),'REG_IADM':RG('AA')}
for n2,v2 in NAMES.items(): wb.defined_names.add(DefinedName(n2, attr_text=v2))

wb.calculation.fullCalcOnLoad=True
wb.properties.title='Registro de Asistencia'
out=SP+'/Registro_Asistencia_2026-1_V11.xlsx'
wb.save(out); print('guardado:',out)
