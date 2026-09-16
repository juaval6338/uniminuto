# -*- coding: utf-8 -*-
"""Reconstruye Registro_Asistencia optimizado."""
import pickle, sys, datetime
SP='/tmp/claude-0/-home-user-uniminuto/b364fe32-f339-5501-a525-09b94e80e92f/scratchpad'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

D = pickle.load(open(SP+'/data.pkl','rb'))

FIRST, NROWS = 19, 2000           # filas de datos en Registros
LAST = FIRST + NROWS - 1          # 2018
EST_LAST, COL_LAST = 60001, 6001  # topes de las bases (con holgura)
EVA_FIRST, EVA_LAST = 10, 509
SEDES = ['BUG','BVA','CHI','CIN','FLO','IPI','PAS','PER','QUI']
TIPOS = ['Estudiante','Docente','Administrativo','Externo']

# ---------- paleta ----------
NAVY   = '1F3864'; BLUE = '2E5B9A'; LIGHT = 'DDEBF7'
INPUT  = 'FFF2CC'; ORANGE='ED7D31'; SOFT = 'FCE4D6'
GREY   = 'F2F2F2'; WHITE='FFFFFF'; GREEN='E2EFDA'
thin   = Side(style='thin', color='BFBFBF')
BOX    = Border(left=thin,right=thin,top=thin,bottom=thin)

def F(sz=10,b=False,c='000000',i=False): return Font(name='Calibri',size=sz,bold=b,color=c,italic=i)
def Fill(c): return PatternFill('solid', fgColor=c)
def AL(h='center',v='center',wrap=False): return Alignment(horizontal=h,vertical=v,wrap_text=wrap)

def style(ws, ref, font=None, fill=None, align=None, border=None, fmt=None):
    for row in ws[ref]:
        for c in row:
            if font: c.font = font
            if fill: c.fill = fill
            if align: c.alignment = align
            if border: c.border = border
            if fmt: c.number_format = fmt

def band(ws, ref, text, bg=NAVY, fg=WHITE, size=11):
    ws.merge_cells(ref)
    first = ref.split(':')[0]
    ws[first] = text
    style(ws, ref, font=F(size,True,fg), fill=Fill(bg), align=AL('left','center'))

wb = Workbook()
wb.remove(wb.active)

# =====================================================================
# 1. REGISTROS
# =====================================================================
rg = wb.create_sheet('Registros')
rg.sheet_properties.tabColor = NAVY

widths = {'A':7,'B':12,'C':17,'D':8,'E':38,'F':30,'G':34,'H':15,'I':16,'J':17,
          'K':26,'L':24,'M':2.5,'N':32,'O':30,'P':14,'Q':8,'R':26,'S':15}
for k,v in widths.items(): rg.column_dimensions[k].width = v
for i in range(20, 41):   # T..AN auxiliares
    L = CL(i); rg.column_dimensions[L].width = 11; rg.column_dimensions[L].hidden = True

# --- encabezado ---
band(rg,'A1:L1','REGISTRO DE ASISTENCIA — UNIMINUTO', NAVY, WHITE, 14)
rg.row_dimensions[1].height = 26
band(rg,'A2:L2','Escriba la cédula o el ID en la columna «Documento». Los demás datos se completan solos.', BLUE, WHITE, 10)

rg['B3']='Actividad / evento:'; rg.merge_cells('C3:F3')
rg['H3']='Fecha del evento:'
rg['B4']='Sede del evento:'
for r in (3,4):
    for col in ('B','H'):
        c=rg[col+str(r)]
        if c.value: c.font=F(10,True,NAVY); c.alignment=AL('right')
style(rg,'C3:F3',font=F(11,True),fill=Fill(INPUT),align=AL('left'),border=BOX)
style(rg,'I3:I3',font=F(11,True),fill=Fill(INPUT),align=AL('center'),border=BOX,fmt='DD/MM/YYYY')
style(rg,'C4:C4',font=F(11,True),fill=Fill(INPUT),align=AL('center'),border=BOX)

# --- buscador ---
band(rg,'B6:L6','¿NO SABE EL DOCUMENTO?  BUSQUE POR NOMBRE AQUÍ', ORANGE, WHITE, 11)
rg['B7']='Escriba parte del nombre:'; rg['B7'].font=F(10,True,NAVY); rg['B7'].alignment=AL('right')
rg.merge_cells('C7:E7')
style(rg,'C7:E7',font=F(11,True),fill=Fill(INPUT),align=AL('left'),border=BOX)
rg['F7']='Base:'; rg['F7'].font=F(10,True,NAVY); rg['F7'].alignment=AL('right')
rg['G7']='Estudiantes'
style(rg,'G7:G7',font=F(10,True),fill=Fill(INPUT),align=AL('center'),border=BOX)
rg['H7']='Copie el documento y peguelo en la columna "Documento"'
rg['H7'].font=F(9,False,'808080',True); rg['H7'].alignment=AL('left')

hdr8 = ['Documento','ID','Apellidos y Nombres','Sede','Programa / Área','Correo']
for j,t in enumerate(hdr8):
    c = rg.cell(row=8, column=3+j, value=t)
style(rg,'C8:H8',font=F(10,True,WHITE),fill=Fill(BLUE),align=AL('center','center',True),border=BOX)

# La base trae varias filas por persona (una por programa matriculado), por eso
# el buscador descarta documentos repetidos antes de mostrar los resultados.
S_RANGES = {'C':'S_DOC','D':'S_ID','E':'S_NOM','F':'S_SEDE','G':'S_PROG','H':'S_MAIL'}
SLOTS, SHOW = 24, 8
for k in range(SLOTS):
    r = 9+k
    if k==0:
        rg['AJ%d'%r] = '=IF(TRIM($C$7&"")="","",IFERROR(MATCH("*"&TRIM($C$7)&"*",S_NOM,0),""))'
    else:
        rg['AJ%d'%r] = ('=IF($AJ{p}="","",IFERROR($AJ{p}+MATCH("*"&TRIM($C$7)&"*",'
                        'INDEX(S_NOM,$AJ{p}+1):INDEX(S_NOM,ROWS(S_NOM)),0),""))').format(p=r-1)
    rg['AK%d'%r] = '=IF($AJ{r}="","",INDEX(S_DOC,$AJ{r})&"")'.format(r=r)
    rg['AL%d'%r] = ('=IF($AK{r}="","",IF(MATCH($AK{r},$AK$9:$AK${e},0)=ROW()-8,1,0))'
                    ).format(r=r, e=8+SLOTS)
    rg['AM%d'%r] = '=IF($AL{r}=1,COUNTIF($AL$9:$AL{r},1),"")'.format(r=r)
for k in range(SHOW):
    r = 9+k
    rg['AN%d'%r] = ('=IFERROR(INDEX($AJ$9:$AJ${e},MATCH(ROWS($AN$9:AN{r}),$AM$9:$AM${e},0)),"")'
                    ).format(r=r, e=8+SLOTS)
    for col,nm in S_RANGES.items():
        rg['%s%d'%(col,r)] = '=IF($AN{r}="","",INDEX({nm},$AN{r})&"")'.format(r=r,nm=nm)
    style(rg,'C%d:H%d'%(r,r),font=F(10),fill=Fill(GREY if k%2 else WHITE),align=AL('left'),border=BOX)
    style(rg,'C%d:D%d'%(r,r),align=AL('center'))
rg.row_dimensions[8].height = 28

# --- cabecera de la tabla ---
HDR = {'A':'Q-Part','B':'Fecha','C':'Documento\n(C.C. o ID)','D':'Sede','E':'Apellidos y Nombres',
       'F':'Programa / Área','G':'Correo electrónico','H':'Teléfono','I':'Tipo de participante',
       'J':'Estado','K':'Actividad / espacio','L':'Observaciones',
       'N':'Nombre','O':'Correo electrónico','P':'Teléfono','Q':'Sede','R':'Programa / Área','S':'Tipo'}
band(rg,'A17:L17','DATOS DE LA PERSONA  (automáticos desde la base de datos)', NAVY, WHITE, 11)
band(rg,'N17:S17','COMPLETE SOLO SI EL DOCUMENTO NO EXISTE EN LA BASE', ORANGE, WHITE, 11)
for col,t in HDR.items():
    rg[col+'18'] = t
style(rg,'A18:L18',font=F(10,True,WHITE),fill=Fill(BLUE),align=AL('center','center',True),border=BOX)
style(rg,'N18:S18',font=F(10,True,WHITE),fill=Fill(ORANGE),align=AL('center','center',True),border=BOX)
rg.row_dimensions[18].height = 34
for i in range(20,41):
    rg.cell(row=18, column=i, value='auxiliar - no modificar').font = F(8,False,'A6A6A6',True)

# T=_doc U=_fCol V=_fEst X=_clave Y=_primera Z/AA/AB=Est AC/AD/AE=Doc AF/AG/AH=Adm

for r in range(FIRST, LAST+1):
    rg['A%d'%r] = '=IF(OR($X{r}="",$B{r}=""),"",1)'.format(r=r)
    rg['D%d'%r] = ('=IF($X{r}="","",IF($V{r}<>"",INDEX(BD_EST_SEDE,$V{r})&"",'
                   'IF($U{r}<>"",INDEX(BD_COL_SEDE,$U{r})&"",UPPER(TRIM($Q{r}&"")))))').format(r=r)
    rg['E%d'%r] = ('=IF($X{r}="","",IF($V{r}<>"",INDEX(BD_EST_NOM,$V{r})&"",'
                   'IF($U{r}<>"",INDEX(BD_COL_NOM,$U{r})&"",UPPER(TRIM($N{r}&"")))))').format(r=r)
    rg['F%d'%r] = ('=IF($X{r}="","",IF($V{r}<>"",INDEX(BD_EST_PROG,$V{r})&"",'
                   'IF($U{r}<>"",INDEX(BD_COL_PROG,$U{r})&"",UPPER(TRIM($R{r}&"")))))').format(r=r)
    rg['G%d'%r] = ('=IF($X{r}="","",IF(TRIM($O{r}&"")<>"",TRIM($O{r}&""),'
                   'IF($V{r}<>"",INDEX(BD_EST_MAIL,$V{r})&"",'
                   'IF($U{r}<>"",IF(INDEX(BD_COL_MAIL,$U{r})&""<>"",INDEX(BD_COL_MAIL,$U{r})&"",'
                   'INDEX(BD_COL_MAIL2,$U{r})&""),""))))').format(r=r)
    rg['H%d'%r] = ('=IF($X{r}="","",IF(TRIM($P{r}&"")<>"",TRIM($P{r}&""),'
                   'IF($V{r}<>"",IF(INDEX(BD_EST_CEL,$V{r})&""<>"",INDEX(BD_EST_CEL,$V{r})&"",'
                   'IF(INDEX(BD_EST_TEL2,$V{r})&""<>"",INDEX(BD_EST_TEL2,$V{r})&"",'
                   'INDEX(BD_EST_TEL3,$V{r})&"")),'
                   'IF($U{r}<>"",INDEX(BD_COL_TEL,$U{r})&"",""))))').format(r=r)
    rg['I%d'%r] = ('=IF($X{r}="","",IF($V{r}<>"","Estudiante",'
                   'IF($U{r}<>"",IF(INDEX(BD_COL_TIPO,$U{r})&""="DOC","Docente","Administrativo"),'
                   'TRIM($S{r}&""))))').format(r=r)
    rg['J%d'%r] = ('=IF($X{r}="","",IF(OR($U{r}<>"",$V{r}<>""),"En base de datos",'
                   'IF(AND($E{r}<>"",$I{r}<>""),"Externo (manual)","FALTAN DATOS")))').format(r=r)
    # auxiliares
    rg['T%d'%r] = ('=IF(TRIM($C{r}&"")="","",IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('
                   'TRIM($C{r}&"")," ",""),".",""),",",""),TRIM($C{r}&"")))').format(r=r)
    rg['U%d'%r] = ('=IF($T{r}="","",IFERROR(MATCH($T{r},BD_COL_CC,0),'
                   'IFERROR(MATCH($T{r},BD_COL_ID,0),"")))').format(r=r)
    rg['V%d'%r] = ('=IF(OR($T{r}="",$U{r}<>""),"",IFERROR(MATCH($T{r},BD_EST_DOC,0),'
                   'IFERROR(MATCH($T{r},BD_EST_ID,0),"")))').format(r=r)
    rg['X%d'%r] = ('=IF($T{r}<>"",$T{r}&"",IF(TRIM($N{r}&"")<>"",UPPER(TRIM($N{r}&"")),""))').format(r=r)
    rg['Y%d'%r] = '=IF($X{r}="","",IF(MATCH($X{r},$X${f}:$X${l},0)=ROW()-{o},1,0))'.format(r=r,f=FIRST,l=LAST,o=FIRST-1)
    for key,(kc,pc,ic) in {'Estudiante':('Z','AA','AB'),'Docente':('AC','AD','AE'),
                           'Administrativo':('AF','AG','AH')}.items():
        rg['%s%d'%(kc,r)] = '=IF(AND($I{r}="{k}",$F{r}<>""),$F{r}&"","")'.format(r=r,k=key)
        rg['%s%d'%(pc,r)] = ('=IF(${kc}{r}="","",IF(MATCH(${kc}{r},${kc}${f}:${kc}${l},0)=ROW()-{o},'
                             '${kc}{r},""))').format(r=r,kc=kc,f=FIRST,l=LAST,o=FIRST-1)
        rg['%s%d'%(ic,r)] = '=IF(${pc}{r}="","",COUNTIF(${pc}${f}:${pc}{r},"?*"))'.format(r=r,pc=pc,f=FIRST)

# --- estilos de las filas de datos ---
f_auto  = F(10); f_in = F(10,False,'7F6000')
fill_a  = Fill(LIGHT); fill_i = Fill(INPUT); fill_k = Fill(GREEN)
for r in range(FIRST, LAST+1):
    for col in 'ADEFGHIJ':
        c = rg[col+str(r)]; c.font=f_auto; c.fill=fill_a; c.border=BOX
        c.alignment = AL('left') if col in 'EFGJ' else AL('center')
    for col in ('B','C','K','L'):
        c = rg[col+str(r)]; c.font=f_in; c.fill=fill_i; c.border=BOX
        c.alignment = AL('center') if col in 'BC' else AL('left')
    for col in ('N','O','P','Q','R','S'):
        c = rg[col+str(r)]; c.font=f_in; c.fill=Fill(SOFT); c.border=BOX
        c.alignment = AL('center') if col in ('P','Q','S') else AL('left')
    rg['A%d'%r].number_format='0'
    rg['B%d'%r].number_format='DD/MM/YYYY'
    rg['C%d'%r].number_format='0'
    rg['P%d'%r].number_format='@'

# --- migración de los registros existentes ---
old = D['reg']          # índice 0 == fila 1
n = 0
for i in range(FIRST-1, min(len(old), 2000)):
    o = old[i]
    if o[5] in (None,''):        # columna F = documento
        continue
    r = FIRST + n; n += 1
    if r > LAST: break
    rg['B%d'%r] = o[4]                                   # E  fecha
    rg['C%d'%r] = o[5]                                   # F  documento
    if o[16] not in (None,''): rg['K%d'%r] = o[16]       # Q  actividad
    if o[11] not in (None,''): rg['O%d'%r] = o[11]       # L  correo manual
    if o[12] not in (None,''): rg['S%d'%r] = o[12]       # M  tipo manual
    if o[13] not in (None,''): rg['Q%d'%r] = o[13]       # N  sede manual
    if o[14] not in (None,''): rg['R%d'%r] = o[14]       # O  programa manual
print('registros migrados:', n)

# --- validaciones ---
def dv(ws, formula, cells, title=None, prompt=None, style='stop'):
    v = DataValidation(type='list', formula1=formula, allow_blank=True, showDropDown=False)
    v.errorStyle = style
    if prompt: v.prompt = prompt; v.promptTitle = title or ''; v.showInputMessage = True
    v.showErrorMessage = (style != 'information')
    ws.add_data_validation(v)
    for c in cells: v.add(c)
    return v

dv(rg, '"%s"'%','.join(SEDES), ['Q%d:Q%d'%(FIRST,LAST), 'C4'])
dv(rg, '"%s"'%','.join(TIPOS), ['S%d:S%d'%(FIRST,LAST)])
dv(rg, '"Estudiantes,Colaboradores"', ['G7'])

vfecha = DataValidation(type='date', operator='greaterThan', formula1='DATE(2015,1,1)',
                        allow_blank=True, errorStyle='warning')
vfecha.promptTitle='Fecha de participación'; vfecha.prompt='Escriba la fecha en formato DD/MM/AAAA.'
vfecha.showInputMessage=True; vfecha.showErrorMessage=True
rg.add_data_validation(vfecha); vfecha.add('B%d:B%d'%(FIRST,LAST)); vfecha.add('I3')

# --- formato condicional: persona no encontrada ---
rg.conditional_formatting.add('A%d:L%d'%(FIRST,LAST),
    FormulaRule(formula=['AND($T%d<>"",$U%d="",$V%d="")'%(FIRST,FIRST,FIRST)],
                fill=Fill('FFE0CC'), stopIfTrue=False))
rg.conditional_formatting.add('N%d:S%d'%(FIRST,LAST),
    FormulaRule(formula=['AND($T%d<>"",$U%d="",$V%d="")'%(FIRST,FIRST,FIRST)],
                fill=Fill('FFF6E5'), stopIfTrue=False))

rg.conditional_formatting.add('J%d:J%d'%(FIRST,LAST),
    FormulaRule(formula=['$J%d="FALTAN DATOS"'%FIRST],
                fill=Fill('FFC7CE'), font=Font(color='9C0006', bold=True), stopIfTrue=True))
rg.conditional_formatting.add('J%d:J%d'%(FIRST,LAST),
    FormulaRule(formula=['$J%d<>""'%FIRST],
                font=Font(color='375623'), stopIfTrue=False))
rg.freeze_panes = 'D19'
rg.auto_filter.ref = 'A18:L%d'%LAST
rg['C3'].comment = Comment('Nombre de la actividad o espacio. Se copia a la columna K de cada fila si usted lo desea.','Ayuda')
rg.sheet_view.showGridLines = False

# =====================================================================
# 2. RESULTADOS POR TIPO
# =====================================================================
rt = wb.create_sheet('Resultados por tipo')
rt.sheet_properties.tabColor = BLUE
rt.sheet_view.showGridLines = False
for k,v in {'A':2,'B':26,'C':16,'D':12,'E':3,'F':26,'G':16,'H':12}.items():
    rt.column_dimensions[k].width = v

band(rt,'B1:H1','RESULTADOS DE LA ACTIVIDAD POR TIPO DE PARTICIPANTE', NAVY, WHITE, 14)
rt.row_dimensions[1].height = 26
rt['B3']='Seleccione Sede:'; rt['B3'].font=F(11,True,NAVY); rt['B3'].alignment=AL('right')
rt['C3']='TODAS'
style(rt,'C3:C3',font=F(12,True,'7F6000'),fill=Fill(INPUT),align=AL('center'),border=BOX)
dv(rt,'"TODAS,%s"'%','.join(SEDES), ['C3'])
rt['D3']='(filtro)'; rt['D3'].font=F(9,False,'808080',True); rt['D3'].alignment=AL('left')

REG = {'qp':'REG_QPART','prim':'REG_PRIMERA','tipo':'REG_TIPO','sede':'REG_SEDE','prog':'REG_PROG'}

def bloque(ws, top, titulo, etiqueta, criterio):
    band(ws,'B%d:D%d'%(top,top), titulo, BLUE, WHITE, 11)
    for j,t in enumerate(['Tipo', etiqueta, '%']):
        ws.cell(row=top+1, column=2+j, value=t)
    style(ws,'B%d:D%d'%(top+1,top+1),font=F(10,True,WHITE),fill=Fill('4472C4'),align=AL('center'),border=BOX)
    filas = [('Estudiantes','Estudiante'),('Administrativos','Administrativo'),
             ('Docentes','Docente'),('Externos','Externo')]
    for k,(lab,val) in enumerate(filas):
        r = top+2+k
        ws['B%d'%r] = lab
        ws['C%d'%r] = ('=IF($C$3="TODAS",COUNTIFS({c},1,{t},"{v}"),'
                       'COUNTIFS({c},1,{t},"{v}",{s},$C$3))').format(c=criterio,t=REG['tipo'],v=val,s=REG['sede'])
        ws['D%d'%r] = '=IFERROR($C{r}/$C${tt},"")'.format(r=r, tt=top+6)
    rtot = top+6
    ws['B%d'%rtot]='Total'
    ws['C%d'%rtot]='=SUM(C%d:C%d)'%(top+2,top+5)
    ws['D%d'%rtot]='=IF($C$%d=0,"",SUM(D%d:D%d))'%(rtot,top+2,top+5)
    style(ws,'B%d:D%d'%(top+2,top+5),font=F(11),fill=Fill(WHITE),align=AL('center'),border=BOX)
    style(ws,'B%d:B%d'%(top+2,top+5),align=AL('left'))
    style(ws,'B%d:D%d'%(rtot,rtot),font=F(11,True),fill=Fill(LIGHT),align=AL('center'),border=BOX)
    style(ws,'B%d:B%d'%(rtot,rtot),align=AL('left'))
    style(ws,'D%d:D%d'%(top+2,rtot),fmt='0.0%')
    style(ws,'C%d:C%d'%(top+2,rtot),fmt='#,##0')

bloque(rt, 5,  'PARTICIPANTES  (personas distintas)',      'Participantes',  REG['prim'])
bloque(rt, 14, 'PARTICIPACIONES  (asistencias registradas)','Participaciones', REG['qp'])
rt['B22']='Las cifras responden al filtro de sede. «TODAS» incluye todas las sedes.'
rt['B22'].font=F(9,False,'808080',True)

# =====================================================================
# 3. RESULTADOS POR PROGRAMA
# =====================================================================
rp = wb.create_sheet('Resultados por programa')
rp.sheet_properties.tabColor = BLUE
rp.sheet_view.showGridLines = False
for k,v in {'A':2,'B':30,'C':15,'D':16,'E':3,'F':30,'G':15,'H':16,'I':3,'J':30,'K':15,'L':16}.items():
    rp.column_dimensions[k].width = v

band(rp,'B1:L1','RESULTADOS DE LA ACTIVIDAD POR PROGRAMA', NAVY, WHITE, 14)
rp.row_dimensions[1].height = 26
rp['B3']='Sede seleccionada:'; rp['B3'].font=F(11,True,NAVY); rp['B3'].alignment=AL('right')
rp['C3']="='Resultados por tipo'!$C$3"
style(rp,'C3:C3',font=F(12,True,NAVY),fill=Fill(LIGHT),align=AL('center'),border=BOX)
rp['D3']='(se cambia en la hoja Resultados por tipo)'
rp['D3'].font=F(9,False,'808080',True); rp['D3'].alignment=AL('left')

PROG_TOP, PROG_N = 7, 60
for cini, titulo, tipo, pn, inx in [(2,'ESTUDIANTES PARTICIPANTES','Estudiante','REG_PROGEST','REG_IEST'),
                                    (6,'DOCENTES PARTICIPANTES','Docente','REG_PROGDOC','REG_IDOC'),
                                    (10,'ADMINISTRATIVOS PARTICIPANTES','Administrativo','REG_PROGADM','REG_IADM')]:
    c0,c1,c2 = CL(cini), CL(cini+1), CL(cini+2)
    band(rp,'%s5:%s5'%(c0,c2), titulo, BLUE, WHITE, 11)
    for j,t in enumerate(['Programa / Área','Participantes','Participaciones']):
        rp.cell(row=PROG_TOP, column=cini+j, value=t)
    style(rp,'%s%d:%s%d'%(c0,PROG_TOP,c2,PROG_TOP),font=F(10,True,WHITE),fill=Fill('4472C4'),
          align=AL('center','center',True),border=BOX)
    for k in range(PROG_N):
        r = PROG_TOP+1+k
        rp['%s%d'%(c0,r)] = ('=IFERROR(INDEX({pn},MATCH(ROWS(${c0}${t}:{c0}{r}),{inx},0))&"","")'
                             ).format(pn=pn,inx=inx,c0=c0,t=PROG_TOP+1,r=r)
        for cc,crit in ((c1,'REG_PRIMERA'),(c2,'REG_QPART')):
            rp['%s%d'%(cc,r)] = ('=IF(${c0}{r}="","",IF($C$3="TODAS",'
                                 'COUNTIFS({cr},1,REG_TIPO,"{tp}",REG_PROG,${c0}{r}),'
                                 'COUNTIFS({cr},1,REG_TIPO,"{tp}",REG_PROG,${c0}{r},REG_SEDE,$C$3)))'
                                 ).format(c0=c0,r=r,cr=crit,tp=tipo)
    rend = PROG_TOP+PROG_N
    style(rp,'%s%d:%s%d'%(c0,PROG_TOP+1,c2,rend),font=F(10),align=AL('center'),border=BOX)
    style(rp,'%s%d:%s%d'%(c0,PROG_TOP+1,c0,rend),align=AL('left'))
    style(rp,'%s%d:%s%d'%(c1,PROG_TOP+1,c2,rend),fmt='#,##0')
rp.freeze_panes = 'A%d'%(PROG_TOP+1)

# =====================================================================
# 4. EVALUACIÓN
# =====================================================================
ev = wb.create_sheet('Evaluación')
ev.sheet_properties.tabColor = '548235'
ev.sheet_view.showGridLines = False
for k,v in {'A':3,'L':3,'M':9,'N':7,'O':4,'S':4,'Y':4,'AB':13,'AC':3}.items():
    ev.column_dimensions[k].width = v
for c in 'BCDEFGHIJK': ev.column_dimensions[c].width = 11
for c in ['P','Q','R','T','U','V','W','X','Z','AA']: ev.column_dimensions[c].width = 6

band(ev,'A1:AB1','REGISTRO DE EVALUACIONES', NAVY, WHITE, 14)
ev.row_dimensions[1].height = 26
ev.merge_cells('O3:AA4')
ev['O3']=('E - Excelente (4)    N - Notable (3)    A - Aceptable (2)    '
          'N/M - Necesita Mejoramiento (1)    N/A - No Aplica (0)')
style(ev,'O3:AA4',font=F(10,True,NAVY),fill=Fill(GREEN),align=AL('center','center',True),border=BOX)

ev.merge_cells('M5:N5'); ev['M5']='Promedio'
ev.merge_cells('M6:N6'); ev['M6']='%'
style(ev,'M5:N6',font=F(10,True,WHITE),fill=Fill(BLUE),align=AL('center'),border=BOX)
ev['AB5']='Promedio %'; style(ev,'AB5:AB5',font=F(10,True,WHITE),fill=Fill(BLUE),align=AL('center',wrap=True),border=BOX)

MIRROR = {'B':'P','C':'Q','D':'R','E':'T','F':'U','G':'V','H':'W','I':'X','J':'Z','K':'AA'}
for src,dst in MIRROR.items():
    ev['%s5'%dst] = '=IFERROR(AVERAGEIF({d}{f}:{d}{l},">0"),"")'.format(d=dst,f=EVA_FIRST,l=EVA_LAST)
    ev['%s6'%dst] = '=IF({d}5<>"",{d}5*0.25,"")'.format(d=dst)
ev['AB6'] = '=IFERROR(AVERAGE(P6:R6,T6:X6,Z6:AA6),"")'
style(ev,'P5:AA6',font=F(9),fill=Fill(LIGHT),align=AL('center'),border=BOX)
style(ev,'P5:AA5',fmt='0.00'); style(ev,'P6:AA6',fmt='0%')
style(ev,'AB6:AB6',font=F(12,True,NAVY),fill=Fill(GREEN),align=AL('center'),border=BOX,fmt='0.0%')

ev['B6']='Ingrese aquí los resultados de las evaluaciones  (puede copiar y pegar)'
style(ev,'B6:K6',font=F(10,True,'7F6000'),fill=Fill(INPUT),align=AL('left'))
ev.merge_cells('B6:K6')

for ref,txt in [('B8:D8','I. Planeación'),('E8:I8','II. Desarrollo'),('J8:K8','III. Aportes'),
                ('P8:R8','I. Planeación'),('T8:X8','II. Desarrollo'),('Z8:AA8','III. Aportes')]:
    band(ev, ref, txt, '4472C4', WHITE, 10)
    style(ev, ref, align=AL('center'))

for j in range(10):
    ev.cell(row=9, column=2+j, value=j+1)
ev['N9']='Nº'; ev['O9']='I'; ev['S9']='II'; ev['Y9']='III'
for j,dst in enumerate(['P','Q','R','T','U','V','W','X','Z','AA']):
    ev['%s9'%dst] = j+1
style(ev,'B9:K9',font=F(10,True,WHITE),fill=Fill(BLUE),align=AL('center'),border=BOX)
style(ev,'N9:AA9',font=F(9,True,WHITE),fill=Fill(BLUE),align=AL('center'),border=BOX)

CODES = '{"0","1","2","3","4","N/A","NA","N/M","NM","A","N","E"}'
for r in range(EVA_FIRST, EVA_LAST+1):
    ev['N%d'%r] = r - EVA_FIRST + 1
    for src,dst in MIRROR.items():
        ev['%s%d'%(dst,r)] = (
            '=IFERROR(IF(TRIM(${s}{r}&"")="","",'
            'IF(ISNUMBER(--TRIM(${s}{r}&"")),'
            'IF(AND(--TRIM(${s}{r}&"")>=0,--TRIM(${s}{r}&"")<=4),--TRIM(${s}{r}&""),"-"),'
            'CHOOSE(MATCH(UPPER(TRIM(${s}{r}&"")),{{"N/A","NA","N/M","NM","A","N","E"}},0),0,0,1,1,2,3,4))),"-")'
        ).format(s=src, r=r)
style(ev,'B%d:K%d'%(EVA_FIRST,EVA_LAST),font=F(11),fill=Fill(INPUT),align=AL('center'),border=BOX)
style(ev,'N%d:N%d'%(EVA_FIRST,EVA_LAST),font=F(9),fill=Fill(GREY),align=AL('center'),border=BOX)
style(ev,'P%d:AA%d'%(EVA_FIRST,EVA_LAST),font=F(9),fill=Fill(WHITE),align=AL('center'),border=BOX)

# datos existentes
for i,row in enumerate(D['eva']):
    if i < 9: continue
    r = EVA_FIRST + (i-9)
    if r > EVA_LAST: break
    for j,col in enumerate('BCDEFGHIJK'):
        v = row[1+j] if 1+j < len(row) else None
        if v not in (None,''): ev['%s%d'%(col,r)] = v

vev = DataValidation(type='list', formula1='"E,N,A,N/M,N/A,4,3,2,1,0"',
                     allow_blank=True, showDropDown=False, errorStyle='warning')
vev.promptTitle='Calificación'; vev.showInputMessage=True
vev.prompt='E=4  N=3  A=2  N/M=1  N/A=0. También puede pegar valores numéricos.'
vev.showErrorMessage=True
ev.add_data_validation(vev); vev.add('B%d:K%d'%(EVA_FIRST,EVA_LAST))

ev.conditional_formatting.add('B%d:K%d'%(EVA_FIRST,EVA_LAST),
    FormulaRule(formula=['AND(TRIM(B%d&"")<>"",ISERROR(MATCH(UPPER(TRIM(B%d&"")),%s,0)))'
                         %(EVA_FIRST,EVA_FIRST,CODES)],
                fill=Fill('FFC7CE'), font=Font(color='9C0006', bold=True), stopIfTrue=False))
ev.freeze_panes = 'B10'
ev['M8']='Valores'; ev['M8'].font=F(9,False,'808080',True)

# =====================================================================
# 5. GRÁFICOS EVALUACIÓN
# =====================================================================
gr = wb.create_sheet('Gráficos Evaluación')
gr.sheet_properties.tabColor = '548235'
gr.sheet_view.showGridLines = False
for c in 'BCDFGHIJLMO': gr.column_dimensions[c].width = 13
for c in ('A','E','K','N'): gr.column_dimensions[c].width = 2.5

band(gr,'A1:O1','GRÁFICAS DE EVALUACIÓN', NAVY, WHITE, 14)
gr.row_dimensions[1].height = 26
PREG = {
 'B':'1. La divulgación de la actividad, evento o servicio.',
 'C':'2. La actividad, evento o servicio se realizó de acuerdo con los tiempos establecidos.',
 'D':'3. Utilización de los recursos tecnológicos y audiovisuales.',
 'F':'4. La claridad en la temática o propósito de la actividad, evento o servicio.',
 'G':'5. La actividad, evento o servicio promovió la participación activa de los asistentes.',
 'H':'6. Actitud y disponibilidad del facilitador o colaborador.',
 'I':'7. El manejo del tema por parte del facilitador o colaborador.',
 'J':'8. El lugar y sus condiciones físicas.',
 'L':'9. Se cumplió su expectativa frente a la actividad, evento o servicio.',
 'M':'10. La actividad, evento o servicio realizado aportó a su formación integral.',
 'O':'TOTAL'}
SRC = {'B':'P6','C':'Q6','D':'R6','F':'T6','G':'U6','H':'V6','I':'W6','J':'X6',
       'L':'Z6','M':'AA6','O':'AB6'}
for ref,txt in [('B3:D3','I. Planeación'),('F3:J3','II. Desarrollo'),
                ('L3:M3','III. Aportes'),('O3:O3','Promedio %')]:
    band(gr, ref, txt, BLUE, WHITE, 11); style(gr, ref, align=AL('center'))
for col,txt in PREG.items():
    gr[col+'4'] = txt
    gr[col+'5'] = "=Evaluación!%s"%SRC[col]
style(gr,'B4:O4',font=F(9),fill=Fill(GREY),align=AL('center','top',True),border=BOX)
style(gr,'B5:O5',font=F(11,True,NAVY),fill=Fill(LIGHT),align=AL('center'),border=BOX,fmt='0.0%')
gr.row_dimensions[4].height = 60

def grafico(titulo, c0, c1, ancla, ancho):
    ch = BarChart(); ch.type='col'; ch.style=10; ch.title=titulo
    ch.y_axis.numFmt='0%'; ch.y_axis.scaling.min=0; ch.y_axis.scaling.max=1
    ch.y_axis.majorGridlines=None; ch.legend=None
    ch.add_data(Reference(gr,min_col=c0,max_col=c1,min_row=5,max_row=5), from_rows=True)
    ch.set_categories(Reference(gr,min_col=c0,max_col=c1,min_row=4,max_row=4))
    ch.dataLabels = __import__('openpyxl.chart.label',fromlist=['DataLabelList']).DataLabelList()
    ch.dataLabels.showVal=True
    ch.width=ancho; ch.height=9
    gr.add_chart(ch, ancla)

grafico('I. Planeación',   2, 4, 'B7',  9)
grafico('II. Desarrollo',  6,10, 'F7', 14)
grafico('III. Aportes',   12,13, 'L7',  7)
grafico('Promedio total', 15,15, 'O7',  6)

# =====================================================================
# 6-7. BASES DE DATOS
# =====================================================================
def base(nombre, filas, color):
    ws = wb.create_sheet(nombre)
    ws.sheet_properties.tabColor = color
    head = filas[0]
    for fila in filas: ws.append(fila)
    ncol = len(head)
    style(ws,'A1:%s1'%CL(ncol), font=F(10,True,WHITE), fill=Fill('595959'),
          align=AL('center','center',True), border=BOX)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes='A2'
    ws.auto_filter.ref='A1:%s%d'%(CL(ncol), len(filas))
    for j,h in enumerate(head):
        L=CL(j+1)
        ws.column_dimensions[L].width = 30 if h and ('NOMBRE' in str(h).upper() or 'CORREO' in str(h).upper()
                                                     or 'PROGRAMA' in str(h).upper()) else 14
        if h and 'FECHA' in str(h).upper():
            for r in range(2, len(filas)+1):
                ws.cell(row=r, column=j+1).number_format='DD/MM/YYYY'
    return ws

base('BD ADM-DOC', D['adm'], '808080')
base('BD EST',     D['est'], '808080')

# =====================================================================
# NOMBRES DEFINIDOS
# =====================================================================
EST = lambda c: "'BD EST'!${c}$2:${c}${n}".format(c=c, n=EST_LAST)
COL = lambda c: "'BD ADM-DOC'!${c}$2:${c}${n}".format(c=c, n=COL_LAST)
RG  = lambda c: "Registros!${c}${f}:${c}${l}".format(c=c, f=FIRST, l=LAST)
NAMES = {
 'BD_EST_ID':EST('B'), 'BD_EST_NOM':EST('C'), 'BD_EST_SEDE':EST('G'), 'BD_EST_PROG':EST('K'),
 'BD_EST_CEL':EST('V'),'BD_EST_TEL2':EST('W'),'BD_EST_TEL3':EST('X'),'BD_EST_MAIL':EST('Y'),
 'BD_EST_DOC':EST('AB'),
 'BD_COL_CC':COL('B'), 'BD_COL_ID':COL('C'), 'BD_COL_NOM':COL('D'), 'BD_COL_TIPO':COL('J'),
 'BD_COL_PROG':COL('K'),'BD_COL_MAIL':COL('L'),'BD_COL_MAIL2':COL('M'),'BD_COL_TEL':COL('N'),
 'BD_COL_SEDE':COL('P'),
 'REG_QPART':RG('A'), 'REG_SEDE':RG('D'), 'REG_PROG':RG('F'), 'REG_TIPO':RG('I'),
 'REG_PRIMERA':RG('Y'),
 'REG_PROGEST':RG('AA'), 'REG_IEST':RG('AB'),
 'REG_PROGDOC':RG('AD'), 'REG_IDOC':RG('AE'),
 'REG_PROGADM':RG('AG'), 'REG_IADM':RG('AH'),
 'S_NOM':'CHOOSE(IF(Registros!$G$7="Colaboradores",2,1),BD_EST_NOM,BD_COL_NOM)',
 'S_DOC':'CHOOSE(IF(Registros!$G$7="Colaboradores",2,1),BD_EST_DOC,BD_COL_CC)',
 'S_ID' :'CHOOSE(IF(Registros!$G$7="Colaboradores",2,1),BD_EST_ID,BD_COL_ID)',
 'S_SEDE':'CHOOSE(IF(Registros!$G$7="Colaboradores",2,1),BD_EST_SEDE,BD_COL_SEDE)',
 'S_PROG':'CHOOSE(IF(Registros!$G$7="Colaboradores",2,1),BD_EST_PROG,BD_COL_PROG)',
 'S_MAIL':'CHOOSE(IF(Registros!$G$7="Colaboradores",2,1),BD_EST_MAIL,BD_COL_MAIL)',
}
for n,v in NAMES.items():
    wb.defined_names.add(DefinedName(n, attr_text=v))

wb.calculation.fullCalcOnLoad = True
wb.properties.title = 'Registro de Asistencia'
wb.properties.creator = 'Bienestar Institucional'
out = SP+'/Registro_Asistencia_2026-1_V10.xlsx'
wb.save(out)
print('guardado:', out)
