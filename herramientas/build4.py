# -*- coding: utf-8 -*-
"""Registro de Asistencia V13."""
import pickle
SP='/tmp/claude-0/-home-user-uniminuto/b364fe32-f339-5501-a525-09b94e80e92f/scratchpad'
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.data_source import AxDataSource, StrRef

D=pickle.load(open(SP+'/data.pkl','rb'))
FIRST,NROWS = 14,1000
LAST=FIRST+NROWS-1                      # 1013
EST_LAST,COL_LAST = 70001,8001          # holgura para alimentar las bases
EVA_FIRST,EVA_LAST = 11,310
SEDES=['BUG','BVA','CHI','CIN','FLO','IPI','PAS','PER','QUI']
TIPOS=['ESTUDIANTE','PROFESOR','ADMINISTRATIVO','EXTERNO']
SLOTS,SHOW = 30,15

AZUL='1F3864'; AZUL2='2E5B9A'; EDIT='FFF4CC'; AUTO='FFFFFF'; TOT='EDEDED'
REPE='FFC000'; ALERTA='FFD9CC'; VERDE='E2EFDA'
thin=Side(style='thin',color='C9C9C9'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def Fo(sz=10,b=False,c='000000',i=False): return Font(name='Calibri',size=sz,bold=b,color=c,italic=i)
def Fi(c): return PatternFill('solid',fgColor='FF'+c)
def Al(h='center',v='center',w=False): return Alignment(horizontal=h,vertical=v,wrap_text=w)
def style(ws,ref,font=None,fill=None,align=None,border=None,fmt=None):
    for row in ws[ref]:
        for c in row:
            if font:c.font=font
            if fill:c.fill=fill
            if align:c.alignment=align
            if border:c.border=border
            if fmt:c.number_format=fmt
def band(ws,ref,text,bg=AZUL,fg='FFFFFF',size=11,h='left'):
    ws.merge_cells(ref); ws[ref.split(':')[0]]=text
    style(ws,ref,font=Fo(size,True,fg),fill=Fi(bg),align=Al(h,'center'))
def lbl(ws,ref,text):
    if ref.split(':')[0]!=ref.split(':')[1]: ws.merge_cells(ref)
    ws[ref.split(':')[0]]=text
    style(ws,ref,font=Fo(10,True,AZUL),align=Al('right','center'))
def dv(ws,f1,cells,tipo='list',**kw):
    v=DataValidation(type=tipo,formula1=f1,allow_blank=True,**kw)   # sin "=" inicial
    ws.add_data_validation(v)
    for c in cells: v.add(c)
    return v

wb=Workbook(); wb.remove(wb.active)

# ============================= REGISTROS =============================
rg=wb.create_sheet('Registros'); rg.sheet_properties.tabColor='FF'+AZUL
rg.sheet_view.showGridLines=False
ANCHOS={'A':7,'B':11,'C':15,'D':10,'E':8,'F':34,'G':28,'H':31,'I':27,'J':13,'K':13,'L':17,
        'M':27,'N':17,'O':9,'P':24,'Q':22,'R':22}
for k,v in ANCHOS.items(): rg.column_dimensions[k].width=v
rg.column_dimensions['A'].hidden=True                 # Q-Part oculta
for i in range(20,42):                                # T..AO auxiliares
    L=CL(i); rg.column_dimensions[L].width=12; rg.column_dimensions[L].hidden=True

band(rg,'B1:R1','REGISTRO DE ASISTENCIA — BIENESTAR INSTITUCIONAL',AZUL,'FFFFFF',14)
rg.row_dimensions[1].height=26
rg.merge_cells('B2:R2')
rg['B2']='Escriba la cédula o el ID en «Documento» y el archivo trae los datos del participante.'
style(rg,'B2:R2',font=Fo(10,True,AZUL2),fill=Fi('E9EFF7'),align=Al('left'))

# ---- buscador ----
band(rg,'B4:G4','BUSCAR PARTICIPANTE POR NOMBRE',AZUL2,'FFFFFF',11)
lbl(rg,'B5:C5','Escriba parte del nombre:')
rg.merge_cells('D5:E5'); style(rg,'D5:E5',font=Fo(11,True,'7F6000'),fill=Fi(EDIT),align=Al('left'),border=BOX)
rg['F5']='Estudiantes'
rg.merge_cells('F5:G5'); style(rg,'F5:G5',font=Fo(10,True,'7F6000'),fill=Fi(EDIT),align=Al('center'),border=BOX)
dv(rg,'"Estudiantes,Colaboradores"',['F5'])
lbl(rg,'B6:C6','Seleccione el nombre:')
rg.merge_cells('D6:G6'); style(rg,'D6:G6',font=Fo(11,True,'7F6000'),fill=Fi(EDIT),align=Al('left'),border=BOX)
dv(rg,'$AM$5:$AM$%d'%(4+SHOW),['D6'])
for r,t in ((7,'Documento:'),(8,'Sede y programa:'),(9,'Correo institucional:')):
    lbl(rg,'B%d:C%d'%(r,r),t)
    rg.merge_cells('D%d:G%d'%(r,r))
    style(rg,'D%d:G%d'%(r,r),font=Fo(11,True,AZUL),fill=Fi(TOT),align=Al('left'),border=BOX)
rg['AO5']='=IFERROR(INDEX($AN$5:$AN$%d,MATCH($D$6,$AM$5:$AM$%d,0)),"")'%(4+SHOW,4+SHOW)
rg['D7']='=IF($AO$5="","",INDEX(S_DOC,$AO$5)&"")'
rg['D8']='=IF($AO$5="","",INDEX(S_SEDE,$AO$5)&"   —   "&INDEX(S_PROG,$AO$5))'
rg['D9']='=IF($AO$5="","",INDEX(S_MAIL,$AO$5)&"")'
for k in range(SLOTS):
    r=5+k
    if k==0:
        rg['AH%d'%r]='=IF(TRIM($D$5&"")="","",IFERROR(MATCH("*"&TRIM($D$5)&"*",S_NOM,0),""))'
    else:
        rg['AH%d'%r]=('=IF($AH{p}="","",IFERROR($AH{p}+MATCH("*"&TRIM($D$5)&"*",'
                      'INDEX(S_NOM,$AH{p}+1):INDEX(S_NOM,ROWS(S_NOM)),0),""))').format(p=r-1)
    rg['AI%d'%r]='=IF($AH{r}="","",INDEX(S_DOC,$AH{r})&"")'.format(r=r)
    rg['AJ%d'%r]='=IF($AI{r}="","",IF(MATCH($AI{r},$AI$5:$AI${e},0)=ROW()-4,1,0))'.format(r=r,e=4+SLOTS)
    rg['AK%d'%r]='=IF($AJ{r}=1,COUNTIF($AJ$5:$AJ{r},1),"")'.format(r=r)
for k in range(SHOW):
    r=5+k
    rg['AN%d'%r]=('=IFERROR(INDEX($AH$5:$AH${e},MATCH(ROWS($AN$5:AN{r}),$AK$5:$AK${e},0)),"")'
                  ).format(r=r,e=4+SLOTS)
    rg['AM%d'%r]='=IF($AN{r}="","",INDEX(S_NOM,$AN{r})&"   —   "&INDEX(S_DOC,$AN{r}))'.format(r=r)

# ---- resumen ----
band(rg,'J4:P4','RESUMEN DE PARTICIPACIÓN',AZUL2,'FFFFFF',11)
lbl(rg,'J5:K5','Filtrar por sede:')
rg['L5']='TODAS'
style(rg,'L5:L5',font=Fo(11,True,'7F6000'),fill=Fi(EDIT),align=Al('center'),border=BOX)
dv(rg,'"TODAS,%s"'%','.join(SEDES),['L5'])
for ref,t in [('J6:K6','Tipo de participante'),('L6:L6','Participantes'),('M6:M6','%'),
              ('N6:O6','Participaciones'),('P6:P6','%')]:
    a,b=ref.split(':')
    if a!=b: rg.merge_cells(ref)
    rg[a]=t
style(rg,'J6:P6',font=Fo(10,True,'FFFFFF'),fill=Fi('4472C4'),align=Al('center','center',True),border=BOX)
for k,t in enumerate(TIPOS):
    r=7+k
    rg.merge_cells('J%d:K%d'%(r,r)); rg['J%d'%r]=t
    rg['L%d'%r]=('=IF($L$5="TODAS",COUNTIFS(REG_PRIM,1,REG_TIPO,"{t}"),'
                 'COUNTIFS(REG_PRIM,1,REG_TIPO,"{t}",REG_SEDE,$L$5))').format(t=t)
    rg['M%d'%r]='=IFERROR($L{r}/$L$11,"")'.format(r=r)
    rg.merge_cells('N%d:O%d'%(r,r))
    rg['N%d'%r]=('=IF($L$5="TODAS",COUNTIFS(REG_QPART,1,REG_TIPO,"{t}"),'
                 'COUNTIFS(REG_QPART,1,REG_TIPO,"{t}",REG_SEDE,$L$5))').format(t=t)
    rg['P%d'%r]='=IFERROR($N{r}/$N$11,"")'.format(r=r)
rg.merge_cells('J11:K11'); rg['J11']='TOTAL'
rg['L11']='=SUM(L7:L10)'; rg['M11']='=IF($L$11=0,"",SUM(M7:M10))'
rg.merge_cells('N11:O11'); rg['N11']='=SUM(N7:N10)'; rg['P11']='=IF($N$11=0,"",SUM(P7:P10))'
style(rg,'J7:P10',font=Fo(10),fill=Fi(AUTO),align=Al('center'),border=BOX)
style(rg,'J7:K10',align=Al('left'))
style(rg,'J11:P11',font=Fo(10,True),fill=Fi(TOT),align=Al('center'),border=BOX)
style(rg,'J11:K11',align=Al('left'))
style(rg,'M7:M11',fmt='0.0%'); style(rg,'P7:P11',fmt='0.0%')
style(rg,'L7:L11',fmt='#,##0'); style(rg,'N7:N11',fmt='#,##0')

# ---- leyenda de colores ----
rg['B12']='Cómo leer los colores:'
style(rg,'B12:B12',font=Fo(9,True,'595959'),align=Al('left'))
rg['C12']='  escribe usted  '
style(rg,'C12:C12',font=Fo(9,True,'7F6000'),fill=Fi(EDIT),align=Al('center'),border=BOX)
rg['D12']='  lo calcula el archivo  '
rg.merge_cells('D12:E12')
style(rg,'D12:E12',font=Fo(9,True,'595959'),fill=Fi(AUTO),align=Al('center'),border=BOX)
rg['F12']='  documento repetido  '
style(rg,'F12:F12',font=Fo(9,True,'7F6000'),fill=Fi(REPE),align=Al('center'),border=BOX)
rg['G12']='  no está en la base  '
style(rg,'G12:G12',font=Fo(9,True,'843C0C'),fill=Fi(ALERTA),align=Al('center'),border=BOX)

# ---- tabla ----
HDR={'A':'Q-Part','B':'Fecha','C':'Documento\n(C.C. o ID)','D':'ID','E':'SEDE',
     'F':'APELLIDOS Y NOMBRES','G':'PROGRAMA / ÁREA','H':'CORREO INSTITUCIONAL\n(@uniminuto.edu)',
     'I':'CORREO ADICIONAL','J':'TELÉFONO','K':'TELÉFONO ADICIONAL','L':'TIPO DE\nPARTICIPANTE',
     'M':'Correo','N':'Tipo','O':'Sede','P':'Programa / Área',
     'Q':'Actividad / espacio','R':'Observaciones'}
band(rg,'M13:P13','COMPLETE SOLO SI EL NOMBRE APARECE COMO «INEXISTENTE»','BF8F00','FFFFFF',9,'center')
for c,t in HDR.items(): rg[c+'14']=t
style(rg,'A14:R14',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center','center',True),border=BOX)
style(rg,'M14:P14',fill=Fi('BF8F00'))
style(rg,'B14:C14',fill=Fi('7F6000'))
rg.row_dimensions[14].height=36
for i in range(20,32):
    rg.cell(row=14,column=i,value='auxiliar - no modificar').font=Fo(8,False,'A6A6A6',True)

F0=FIRST+1                      # primera fila de datos = 15
L0=F0+NROWS-1
for r in range(F0,L0+1):
    rg['A%d'%r]='=IF(OR($T{r}="",$B{r}=""),"",1)'.format(r=r)
    rg['D%d'%r]=('=IF($T{r}="","",IF($V{r}<>"",INDEX(BD_EST_ID,$V{r})&"",'
                 'IF($U{r}<>"",INDEX(BD_COL_ID,$U{r})&"","")))').format(r=r)
    rg['E%d'%r]=('=IF($T{r}="","",UPPER(IF($V{r}<>"",INDEX(BD_EST_SEDE,$V{r})&"",'
                 'IF($U{r}<>"",INDEX(BD_COL_SEDE,$U{r})&"",TRIM($O{r}&"")))))').format(r=r)
    rg['F%d'%r]=('=IF($T{r}="","",UPPER(IF($V{r}<>"",INDEX(BD_EST_NOM,$V{r})&"",'
                 'IF($U{r}<>"",INDEX(BD_COL_NOM,$U{r})&"","INEXISTENTE"))))').format(r=r)
    rg['G%d'%r]=('=IF($T{r}="","",UPPER(IF($V{r}<>"",INDEX(BD_EST_PROG,$V{r})&"",'
                 'IF($U{r}<>"",INDEX(BD_COL_PROG,$U{r})&"",TRIM($P{r}&"")))))').format(r=r)
    rg['H%d'%r]=('=IF($T{r}="","",IF($V{r}<>"",'
                 'IF(INDEX(BD_EST_MAIL,$V{r})&""<>"",INDEX(BD_EST_MAIL,$V{r})&"",TRIM($M{r}&"")),'
                 'IF($U{r}<>"",IF(INDEX(BD_COL_MAIL,$U{r})&""<>"",INDEX(BD_COL_MAIL,$U{r})&"",'
                 'TRIM($M{r}&"")),TRIM($M{r}&""))))').format(r=r)
    rg['I%d'%r]='=IF(OR($Z{r}="",$Z{r}=$H{r}),"",$Z{r})'.format(r=r)
    rg['J%d'%r]=('=IF($T{r}="","",IF($V{r}<>"",INDEX(BD_EST_CEL,$V{r})&"",'
                 'IF($U{r}<>"",INDEX(BD_COL_TEL,$U{r})&"","")))').format(r=r)
    rg['K%d'%r]=('=IF(OR($T{r}="",$V{r}=""),"",'
                 'IF(AND(INDEX(BD_EST_TEL2,$V{r})&""<>"",INDEX(BD_EST_TEL2,$V{r})&""<>$J{r}),'
                 'INDEX(BD_EST_TEL2,$V{r})&"",'
                 'IF(AND(INDEX(BD_EST_TEL3,$V{r})&""<>"",INDEX(BD_EST_TEL3,$V{r})&""<>$J{r}),'
                 'INDEX(BD_EST_TEL3,$V{r})&"","")))').format(r=r)
    rg['L%d'%r]=('=IF($T{r}="","",IF($V{r}<>"",IF(INDEX(BD_EST_COD,$V{r})=1,"ESTUDIANTE",""),'
                 'IF($U{r}<>"",IF(INDEX(BD_COL_COD,$U{r})=27,"PROFESOR",'
                 'IF(INDEX(BD_COL_COD,$U{r})=28,"ADMINISTRATIVO","")),UPPER(TRIM($N{r}&"")))))').format(r=r)
    rg['T%d'%r]=('=IF(TRIM($C{r}&"")="","",IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE('
                 'TRIM($C{r}&"")," ",""),".",""),",",""),TRIM($C{r}&"")))').format(r=r)
    rg['U%d'%r]=('=IF($T{r}="","",IFERROR(MATCH($T{r},BD_COL_CC,0),'
                 'IFERROR(MATCH($T{r},BD_COL_ID,0),"")))').format(r=r)
    rg['V%d'%r]=('=IF(OR($T{r}="",$U{r}<>""),"",IFERROR(MATCH($T{r},BD_EST_DOC,0),'
                 'IFERROR(MATCH($T{r},BD_EST_ID,0),"")))').format(r=r)
    rg['W%d'%r]='=IF($T{r}="","",$T{r}&"")'.format(r=r)
    rg['X%d'%r]=('=IF($W{r}="","",IF(MATCH($W{r},$W${f}:$W${l},0)=ROW()-{o},1,0))'
                 ).format(r=r,f=F0,l=L0,o=F0-1)
    rg['Y%d'%r]='=IF($W{r}="","",IF(COUNTIF($W${f}:$W${l},$W{r})>1,1,0))'.format(r=r,f=F0,l=L0)
    rg['Z%d'%r]=('=IF($T{r}="","",IF($V{r}<>"",IFERROR(LEFT(INDEX(BD_EST_MAIL2,$V{r})&"",'
                 'FIND("#",INDEX(BD_EST_MAIL2,$V{r})&"")-1),INDEX(BD_EST_MAIL2,$V{r})&""),'
                 'IF($U{r}<>"",INDEX(BD_COL_MAIL2,$U{r})&"","")))').format(r=r)
    rg['AA%d'%r]='=IF(AND($L{r}="ESTUDIANTE",$G{r}<>""),$G{r}&"","")'.format(r=r)
    rg['AB%d'%r]=('=IF($AA{r}="","",IF(MATCH($AA{r},$AA${f}:$AA${l},0)=ROW()-{o},$AA{r},""))'
                  ).format(r=r,f=F0,l=L0,o=F0-1)
    rg['AC%d'%r]='=IF($AB{r}="","",COUNTIF($AB${f}:$AB{r},"?*"))'.format(r=r,f=F0)
    rg['AD%d'%r]=('=IF(AND(OR($L{r}="PROFESOR",$L{r}="ADMINISTRATIVO"),$G{r}<>""),$G{r}&"","")'
                  ).format(r=r)
    rg['AE%d'%r]=('=IF($AD{r}="","",IF(MATCH($AD{r},$AD${f}:$AD${l},0)=ROW()-{o},$AD{r},""))'
                  ).format(r=r,f=F0,l=L0,o=F0-1)
    rg['AF%d'%r]='=IF($AE{r}="","",COUNTIF($AE${f}:$AE{r},"?*"))'.format(r=r,f=F0)

f_auto=Fo(10); f_in=Fo(10,False,'7F6000')
AUTOCOLS='ADEFGHIJKL'; EDITCOLS=('B','C','M','N','O','P','Q','R')
for r in range(F0,L0+1):
    for col in AUTOCOLS:
        c=rg[col+str(r)]; c.font=f_auto; c.fill=Fi(AUTO); c.border=BOX
        c.alignment=Al('left') if col in 'FGHI' else Al('center')
    for col in EDITCOLS:
        c=rg[col+str(r)]; c.font=f_in; c.fill=Fi(EDIT); c.border=BOX
        c.alignment=Al('center') if col in ('B','C','N','O') else Al('left')
    rg['A%d'%r].number_format='0'; rg['B%d'%r].number_format='DD/MM/YYYY'
    rg['C%d'%r].number_format='0'

old=D['reg']; n=0
for i in range(18,min(len(old),2000)):
    o=old[i]
    if o[5] in (None,''): continue
    r=F0+n; n+=1
    if r>L0: break
    rg['B%d'%r]=o[4]; rg['C%d'%r]=o[5]
    if o[16] not in (None,''): rg['Q%d'%r]=o[16]
    if o[11] not in (None,''): rg['M%d'%r]=o[11]
    if o[12] not in (None,''): rg['N%d'%r]=str(o[12]).upper()
    if o[13] not in (None,''): rg['O%d'%r]=o[13]
    if o[14] not in (None,''): rg['P%d'%r]=o[14]
print('registros migrados:',n)

dv(rg,'"%s"'%','.join(TIPOS),['N%d:N%d'%(F0,L0)])
dv(rg,'"%s"'%','.join(SEDES),['O%d:O%d'%(F0,L0)])
v=dv(rg,'DATE(2015,1,1)',['B%d:B%d'%(F0,L0)],tipo='date',operator='greaterThan')
v.errorStyle='warning'; v.showErrorMessage=True
# documento repetido -> se resalta el ID
rg.conditional_formatting.add('C%d:D%d'%(F0,L0),
    FormulaRule(formula=['$Y%d=1'%F0],fill=Fi(REPE),font=Font(bold=True,color='7F6000'),stopIfTrue=True))
rg.conditional_formatting.add('B%d:L%d'%(F0,L0),
    FormulaRule(formula=['$F%d="INEXISTENTE"'%F0],fill=Fi(ALERTA),stopIfTrue=False))
rg.freeze_panes='D15'
rg.auto_filter.ref='B14:R%d'%L0

# ====================== RESULTADOS POR PROGRAMA ======================
rp=wb.create_sheet('Resultados por programa'); rp.sheet_properties.tabColor='FF'+AZUL2
rp.sheet_view.showGridLines=False
for k,v in {'A':2,'B':42,'C':16,'D':16,'E':12}.items(): rp.column_dimensions[k].width=v
for i in range(7,16):
    L=CL(i); rp.column_dimensions[L].width=12; rp.column_dimensions[L].hidden=True
band(rp,'B1:E1','PARTICIPACIÓN POR PROGRAMA Y ÁREA',AZUL,'FFFFFF',14)
rp.row_dimensions[1].height=26
lbl(rp,'B3:B3','Sede:')
rp['C3']="=Registros!$L$5"
style(rp,'C3:C3',font=Fo(12,True,AZUL),fill=Fi(TOT),align=Al('center'),border=BOX)
rp.merge_cells('D3:E3'); rp['D3']='(el filtro se cambia en la hoja Registros)'
style(rp,'D3:E3',font=Fo(9,False,'808080',True),align=Al('left'))
NP=40
def bloque_prog(top,titulo,encab,pn,ix,tipos,aux):
    ap,ac,ak,ao=aux
    band(rp,'B%d:E%d'%(top,top),titulo,AZUL2,'FFFFFF',11)
    for j,t in enumerate([encab,'Participantes','Participaciones','% del total']):
        rp.cell(row=top+1,column=2+j,value=t)
    style(rp,'B%d:E%d'%(top+1,top+1),font=Fo(10,True,'FFFFFF'),fill=Fi('4472C4'),
          align=Al('center','center',True),border=BOX)
    h0,h1=top+2,top+1+NP
    cnt=lambda crit,ref:' + '.join(
        ('IF($C$3="TODAS",COUNTIFS({c},1,REG_TIPO,"{t}",REG_PROG,{r}),'
         'COUNTIFS({c},1,REG_TIPO,"{t}",REG_PROG,{r},REG_SEDE,$C$3))').format(c=crit,t=t,r=ref)
        for t in tipos)
    for k in range(NP):
        r=h0+k
        rp['%s%d'%(ap,r)]=('=IFERROR(INDEX({pn},MATCH(ROWS(${a}${h}:{a}{r}),{ix},0))&"","")'
                           ).format(pn=pn,ix=ix,a=ap,h=h0,r=r)
        rp['%s%d'%(ac,r)]='=IF(${a}{r}="","",{f})'.format(a=ap,r=r,f=cnt('REG_PRIM','$%s%d'%(ap,r)))
        rp['%s%d'%(ak,r)]='=IF(${a}{r}="","",${c}{r}*10000+(10000-ROW()))'.format(a=ap,c=ac,r=r)
        rp['%s%d'%(ao,r)]=('=IFERROR(MATCH(LARGE(${k}${h}:${k}${e},ROWS(${o}${h}:{o}{r})),'
                           '${k}${h}:${k}${e},0),"")').format(k=ak,o=ao,h=h0,e=h1,r=r)
        rp['B%d'%r]='=IF(${o}{r}="","",INDEX(${a}${h}:${a}${e},${o}{r})&"")'.format(o=ao,a=ap,h=h0,e=h1,r=r)
        rp['C%d'%r]='=IF(${o}{r}="","",INDEX(${c}${h}:${c}${e},${o}{r}))'.format(o=ao,c=ac,h=h0,e=h1,r=r)
        rp['D%d'%r]='=IF($B{r}="","",{f})'.format(r=r,f=cnt('REG_QPART','$B%d'%r))
        rp['E%d'%r]='=IFERROR($C{r}/$C${t},"")'.format(r=r,t=h1+1)
    rtot=h1+1
    rp['B%d'%rtot]='TOTAL'
    rp['C%d'%rtot]='=SUM(C%d:C%d)'%(h0,h1); rp['D%d'%rtot]='=SUM(D%d:D%d)'%(h0,h1)
    rp['E%d'%rtot]='=IF($C$%d=0,"",SUM(E%d:E%d))'%(rtot,h0,h1)
    style(rp,'B%d:E%d'%(h0,h1),font=Fo(10),fill=Fi(AUTO),align=Al('center'),border=BOX)
    style(rp,'B%d:B%d'%(h0,h1),align=Al('left'))
    style(rp,'B%d:E%d'%(rtot,rtot),font=Fo(10,True),fill=Fi(TOT),align=Al('center'),border=BOX)
    style(rp,'B%d:B%d'%(rtot,rtot),align=Al('left'))
    style(rp,'C%d:D%d'%(h0,rtot),fmt='#,##0'); style(rp,'E%d:E%d'%(h0,rtot),fmt='0.0%')
    return rtot
fin1=bloque_prog(5,'PROGRAMAS ACADÉMICOS  —  estudiantes','Programa académico',
                 'REG_PROGEST','REG_IEST',['ESTUDIANTE'],('G','H','I','J'))
bloque_prog(fin1+3,'ÁREAS Y DEPENDENCIAS  —  profesores y administrativos','Área / dependencia',
            'REG_PROGCOL','REG_ICOL',['PROFESOR','ADMINISTRATIVO'],('K','L','M','N'))
rp.freeze_panes='A7'

# ============================ EVALUACIÓN ============================
PREGUNTAS=[
 ('PLANEACIÓN','Divulgación de la actividad',
  '1. La divulgación de la actividad, evento o servicio fue'),
 ('PLANEACIÓN','Cumplimiento de los tiempos',
  '2. La actividad, evento o servicio se realizó de acuerdo con los tiempos establecidos'),
 ('PLANEACIÓN','Recursos tecnológicos y audiovisuales',
  '3. Los recursos tecnológicos y audiovisuales utilizados fueron'),
 ('DESARROLLO','Claridad de la temática',
  '4. La claridad en la temática o propósito de la actividad, evento o servicio fue'),
 ('DESARROLLO','Participación de los asistentes',
  '5. La actividad, evento o servicio promovió la participación activa de los asistentes'),
 ('DESARROLLO','Actitud del facilitador',
  '6. La Actitud y disponibilidad del facilitador o colaborador fue'),
 ('DESARROLLO','Manejo del tema',
  '7. El manejo del tema por parte del facilitador o colaborador fue'),
 ('DESARROLLO','Lugar y condiciones físicas',
  '8. El lugar y sus condiciones físicas fueron'),
 ('APORTES','Cumplimiento de expectativas',
  '9. Se cumplió su expectativa frente a la actividad, evento o servicio'),
 ('APORTES','Aporte a la formación integral',
  '10. La actividad, evento o servicio realizado aportó a su formación integral'),
]
ENT=['B','C','D','E','F','G','H','I','J','K']
ESP=['P','Q','R','T','U','V','W','X','Z','AA']

ev=wb.create_sheet('Evaluación'); ev.sheet_properties.tabColor='FF548235'
ev.sheet_view.showGridLines=False
for k,v in {'A':3,'L':3,'M':9,'N':7,'O':4,'S':4,'Y':4,'AB':13,'AC':3}.items():
    ev.column_dimensions[k].width=v
for c in ENT: ev.column_dimensions[c].width=13
for c in ESP: ev.column_dimensions[c].width=6
band(ev,'A1:AB1','REGISTRO DE EVALUACIONES',AZUL,'FFFFFF',14)
ev.row_dimensions[1].height=26
ev.merge_cells('A2:AB2')
ev['A2']='Formato FR-BM-DFB-03  ·  Versión 1  ·  Enero 28 de 2021  ·  Plan de Bienestar Institucional'
style(ev,'A2:AB2',font=Fo(9,False,'595959',True),align=Al('left'))
ev.merge_cells('B4:AB5')
ev['B4']=('Para la calificación tenga en cuenta:      '
          'E - Excelente (se superaron mis expectativas) = 4      '
          'N - Notable (se cumplieron mis expectativas) = 3      '
          'A - Aceptable (se cumplieron parcialmente mis expectativas) = 2      '
          'N/M - Necesita Mejoramiento (puede mejorar) = 1      '
          'N/A - No aplica')
style(ev,'B4:AB5',font=Fo(10,True,AZUL),fill=Fi(VERDE),align=Al('left','center',True),border=BOX)
ev.row_dimensions[4].height=18; ev.row_dimensions[5].height=18
ev.merge_cells('M7:N7'); ev['M7']='Promedio'
ev.merge_cells('M8:N8'); ev['M8']='%'
style(ev,'M7:N8',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
ev['AB7']='Promedio %'
style(ev,'AB7:AB7',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center',w=True),border=BOX)
for s,d in zip(ENT,ESP):
    ev['%s7'%d]='=IFERROR(AVERAGEIF({d}{f}:{d}{l},">0"),"")'.format(d=d,f=EVA_FIRST+1,l=EVA_LAST+1)
    ev['%s8'%d]='=IF({d}7<>"",{d}7*0.25,"")'.format(d=d)
ev['AB8']='=IFERROR(AVERAGE(P8:R8,T8:X8,Z8:AA8),"")'
style(ev,'P7:AA8',font=Fo(9),fill=Fi(TOT),align=Al('center'),border=BOX)
style(ev,'P7:AA7',fmt='0.00'); style(ev,'P8:AA8',fmt='0%')
style(ev,'AB8:AB8',font=Fo(12,True,AZUL),fill=Fi(VERDE),align=Al('center'),border=BOX,fmt='0.0%')
for ref,txt in [('B10:D10','I. PLANEACIÓN'),('E10:I10','II. DESARROLLO'),('J10:K10','III. APORTES'),
                ('P10:R10','I. PLANEACIÓN'),('T10:X10','II. DESARROLLO'),('Z10:AA10','III. APORTES')]:
    band(ev,ref,txt,'4472C4','FFFFFF',10,'center')
for j,(sec,corta,larga) in enumerate(PREGUNTAS):
    ev['%s11'%ENT[j]]='%d. %s'%(j+1,corta)
    ev['%s11'%ESP[j]]=j+1
ev['N11']='Nº'; ev['O11']='I'; ev['S11']='II'; ev['Y11']='III'
style(ev,'B11:K11',font=Fo(9,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center','center',True),border=BOX)
style(ev,'N11:AA11',font=Fo(9,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
ev.row_dimensions[11].height=48
ev.merge_cells('B9:K9')
ev['B9']='Ingrese aquí los resultados de las evaluaciones (puede copiar y pegar)'
style(ev,'B9:K9',font=Fo(10,True,'7F6000'),fill=Fi(EDIT),align=Al('left'))
CODES='{"0","1","2","3","4","N/A","NA","N/M","NM","A","N","E"}'
for r in range(EVA_FIRST+1,EVA_LAST+2):
    ev['N%d'%r]=r-EVA_FIRST
    for s,d in zip(ENT,ESP):
        ev['%s%d'%(d,r)]=('=IFERROR(IF(TRIM(${s}{r}&"")="","",'
            'IF(ISNUMBER(--TRIM(${s}{r}&"")),'
            'IF(AND(--TRIM(${s}{r}&"")>=0,--TRIM(${s}{r}&"")<=4),--TRIM(${s}{r}&""),"-"),'
            'CHOOSE(MATCH(UPPER(TRIM(${s}{r}&"")),{{"N/A","NA","N/M","NM","A","N","E"}},0),'
            '0,0,1,1,2,3,4))),"-")').format(s=s,r=r)
style(ev,'B%d:K%d'%(EVA_FIRST+1,EVA_LAST+1),font=Fo(11),fill=Fi(EDIT),align=Al('center'),border=BOX)
style(ev,'N%d:N%d'%(EVA_FIRST+1,EVA_LAST+1),font=Fo(9),fill=Fi(TOT),align=Al('center'),border=BOX)
style(ev,'P%d:AA%d'%(EVA_FIRST+1,EVA_LAST+1),font=Fo(9),fill=Fi(AUTO),align=Al('center'),border=BOX)
for i,row in enumerate(D['eva']):
    if i<9: continue
    r=EVA_FIRST+1+(i-9)
    if r>EVA_LAST+1: break
    for j,col in enumerate(ENT):
        v2=row[1+j] if 1+j<len(row) else None
        if v2 not in (None,''): ev['%s%d'%(col,r)]=v2
vev=dv(ev,'"E,N,A,N/M,N/A,4,3,2,1,0"',['B%d:K%d'%(EVA_FIRST+1,EVA_LAST+1)])
vev.errorStyle='warning'; vev.showErrorMessage=True
vev.promptTitle='Calificación'; vev.showInputMessage=True
vev.prompt='E=4  N=3  A=2  N/M=1  N/A=0. También puede pegar números.'
ev.conditional_formatting.add('B%d:K%d'%(EVA_FIRST+1,EVA_LAST+1),
    FormulaRule(formula=['AND(TRIM(B%d&"")<>"",ISERROR(MATCH(UPPER(TRIM(B%d&"")),%s,0)))'
                         %(EVA_FIRST+1,EVA_FIRST+1,CODES)],
                fill=Fi('FFC7CE'),font=Font(color='9C0006',bold=True),stopIfTrue=False))
ev.freeze_panes='B12'

# ====================== GRÁFICOS EVALUACIÓN ======================
gr=wb.create_sheet('Gráficos Evaluación'); gr.sheet_properties.tabColor='FF548235'
gr.sheet_view.showGridLines=False
for k,v in {'A':2,'B':5,'C':36,'D':11}.items(): gr.column_dimensions[k].width=v
band(gr,'B1:D1','RESULTADOS DE LA EVALUACIÓN',AZUL,'FFFFFF',14)
gr.row_dimensions[1].height=26
gr.merge_cells('B3:C3'); gr['B3']='Aspecto evaluado'; gr['D3']='%'
style(gr,'B3:D3',font=Fo(10,True,'FFFFFF'),fill=Fi(AZUL2),align=Al('center'),border=BOX)
for k,(sec,corta,larga) in enumerate(PREGUNTAS):
    r=4+k
    gr['B%d'%r]=k+1; gr['C%d'%r]=corta; gr['D%d'%r]='=Evaluación!%s8'%ESP[k]
style(gr,'B4:D13',font=Fo(10),fill=Fi(AUTO),align=Al('left'),border=BOX)
style(gr,'B4:B13',align=Al('center'))
style(gr,'D4:D13',font=Fo(10,True,AZUL),fill=Fi(TOT),align=Al('center'),fmt='0.0%')
gr.merge_cells('B15:C15'); gr['B15']='PROMEDIO TOTAL'; gr['D15']='=Evaluación!AB8'
style(gr,'B15:D15',font=Fo(12,True,'FFFFFF'),fill=Fi(AZUL),align=Al('center'),border=BOX)
style(gr,'D15:D15',fmt='0.0%')

ch=BarChart(); ch.type='bar'; ch.title=None; ch.legend=None
ch.y_axis.numFmt='0%'; ch.y_axis.scaling.min=0; ch.y_axis.scaling.max=1   # eje de valores
ch.y_axis.majorGridlines=None; ch.x_axis.majorGridlines=None
ch.x_axis.delete=False; ch.y_axis.delete=False
ch.gapWidth=45
ch.add_data(Reference(gr,min_col=4,min_row=4,max_row=13),titles_from_data=False)
ch.series[0].cat=AxDataSource(strRef=StrRef(f="'Gráficos Evaluación'!$C$4:$C$13"))
ch.series[0].graphicalProperties=GraphicalProperties(solidFill=AZUL2)
dl=DataLabelList()
dl.showVal=True; dl.showSerName=False; dl.showCatName=False
dl.showLegendKey=False; dl.showPercent=False; dl.showBubbleSize=False
dl.numFmt='0.0%'
ch.dataLabels=dl
ch.width=19; ch.height=11
gr.add_chart(ch,'F3')
gr['B18']='Preguntas completas de la encuesta (formato FR-BM-DFB-03)'
gr['B18'].font=Fo(9,True,'595959')
for k,(sec,corta,larga) in enumerate(PREGUNTAS):
    r=19+k
    gr.merge_cells('B%d:D%d'%(r,r)); gr['B%d'%r]=larga
    style(gr,'B%d:D%d'%(r,r),font=Fo(8,False,'808080'),align=Al('left'))

# ============================== BASES ==============================
def base(nombre,filas,nota):
    ws=wb.create_sheet(nombre); ws.sheet_properties.tabColor='FF808080'
    for f in filas: ws.append(f)
    head=filas[0]; ncol=len(head)
    style(ws,'A1:%s1'%CL(ncol),font=Fo(10,True,'FFFFFF'),fill=Fi('595959'),
          align=Al('center','center',True),border=BOX)
    ws.row_dimensions[1].height=30; ws.freeze_panes='A2'
    ws.auto_filter.ref='A1:%s%d'%(CL(ncol),len(filas))
    for j,h in enumerate(head):
        L=CL(j+1); hu=str(h).upper() if h else ''
        ws.column_dimensions[L].width=30 if ('NOMBRE' in hu or 'CORREO' in hu or 'PROGRAMA' in hu) else 14
        if 'FECHA' in hu: ws.column_dimensions[L].number_format='DD/MM/YYYY'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    return ws
base('BD ADM-DOC',D['adm'],'colaboradores')
base('BD EST',D['est'],'estudiantes')

# ========================= NOMBRES DEFINIDOS =========================
EST=lambda c:"'BD EST'!${c}$2:${c}${n}".format(c=c,n=EST_LAST)
COL=lambda c:"'BD ADM-DOC'!${c}$2:${c}${n}".format(c=c,n=COL_LAST)
RG =lambda c:"Registros!${c}${f}:${c}${l}".format(c=c,f=F0,l=L0)
NAMES={
 'BD_EST_ID':EST('B'),'BD_EST_NOM':EST('C'),'BD_EST_SEDE':EST('G'),'BD_EST_PROG':EST('K'),
 'BD_EST_CEL':EST('V'),'BD_EST_TEL2':EST('W'),'BD_EST_TEL3':EST('X'),
 'BD_EST_MAIL':EST('Y'),'BD_EST_MAIL2':EST('Z'),'BD_EST_DOC':EST('AB'),'BD_EST_COD':EST('AU'),
 'BD_COL_CC':COL('B'),'BD_COL_ID':COL('C'),'BD_COL_NOM':COL('D'),'BD_COL_COD':COL('I'),
 'BD_COL_PROG':COL('K'),'BD_COL_MAIL':COL('L'),'BD_COL_MAIL2':COL('M'),'BD_COL_TEL':COL('N'),
 'BD_COL_SEDE':COL('P'),
 'REG_QPART':RG('A'),'REG_SEDE':RG('E'),'REG_PROG':RG('G'),'REG_TIPO':RG('L'),'REG_PRIM':RG('X'),
 'REG_PROGEST':RG('AB'),'REG_IEST':RG('AC'),
 'REG_PROGCOL':RG('AE'),'REG_ICOL':RG('AF'),
 'S_NOM':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_NOM,BD_COL_NOM)',
 'S_DOC':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_DOC,BD_COL_CC)',
 'S_SEDE':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_SEDE,BD_COL_SEDE)',
 'S_PROG':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_PROG,BD_COL_PROG)',
 'S_MAIL':'CHOOSE(IF(Registros!$F$5="Colaboradores",2,1),BD_EST_MAIL,BD_COL_MAIL)'}
for n2,v2 in NAMES.items(): wb.defined_names.add(DefinedName(n2,attr_text=v2))
wb.calculation.fullCalcOnLoad=True
wb.properties.title='Registro de Asistencia'
out=SP+'/_v13_raw.xlsx'
wb.save(out); print('guardado bruto:',out)
