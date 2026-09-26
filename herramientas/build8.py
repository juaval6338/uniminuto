# -*- coding: utf-8 -*-
"""V17: parte del archivo original (V9) y solo cambia el funcionamiento."""
import sys, time
from copy import copy
SP='/tmp/claude-0/-home-user-uniminuto/b364fe32-f339-5501-a525-09b94e80e92f/scratchpad'
sys.path.insert(0,SP)
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.formatting.formatting import ConditionalFormattingList
from sharedstr import convert

t0=time.time()
wb=load_workbook(SP+'/orig/original.xlsx')
print('cargado en %.0fs'%(time.time()-t0))

FIRST, LAST = 19, 2018           # 2000 filas de registro
SEDES=['BUG','BVA','CHI','CIN','FLO','IPI','PAS','PER','QUI']
TIPOS=['ESTUDIANTE','PROFESOR','ADMINISTRATIVO','EXTERNO']
def cp(dst,src): dst._style=copy(src._style)
thin=Side(style='thin',color='FF000000')
BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
INPUT_FILL=PatternFill('solid',fgColor='FFFFFFCC')     # el amarillo claro del original (H8:J8)
NOFILL=PatternFill(fill_type=None)

def dv(ws,f1,cells,tipo='list',**kw):
    v=DataValidation(type=tipo,formula1=f1,allow_blank=True,**kw)
    ws.add_data_validation(v)
    for c in cells: v.add(c)
    return v

# =============================== REGISTROS ===============================
rg=wb['Registros']
# plantillas de estilo tomadas del propio original
ST_AUTO = copy(rg['G19']._style)      # gris, automático
ST_AUTO_L = copy(rg['J19']._style)    # gris, alineado a la izquierda
ST_HDR_AUTO = copy(rg['R18']._style)  # encabezado de columna automática
ST_BUS_HDR = copy(rg['E13']._style)   # encabezado del buscador
ST_BUS_VAL = copy(rg['E14']._style)   # valor gris del buscador
ST_BUS_TIT = copy(rg['E12']._style)   # título del buscador

# filas 2001..2018 heredan el estilo de la fila 2000
for r in range(2001, LAST+1):
    rg.row_dimensions[r].height = rg.row_dimensions[2000].height
    for c in range(1,19):
        cp(rg.cell(row=r,column=c), rg.cell(row=2000,column=c))
# limpiar fórmulas viejas y columnas auxiliares
for r in range(FIRST, rg.max_row+1):
    for c in (1,2,3,7,8,9,10,11,16,18) + tuple(range(19,41)):
        cell=rg.cell(row=r,column=c)
        if c>=19 or (isinstance(cell.value,str) and cell.value.startswith('=')) or r>LAST:
            cell.value=None
# nuevas columnas S y T con el estilo de R; U separador
rg.column_dimensions['S'].width=16; rg.column_dimensions['T'].width=34; rg.column_dimensions['U'].width=4
for c in ('S','T','U','V','W'): rg.column_dimensions[c].hidden=False
cp(rg['S18'],rg['R18']); cp(rg['T18'],rg['R18'])
for c in ('T','U'): cp(rg[c+'10'],rg['S10'])        # la franja gris del título llega hasta la T
for r in range(FIRST,LAST+1):
    cp(rg['S%d'%r],rg['R%d'%r]); cp(rg['T%d'%r],rg['J%d'%r])
# columnas auxiliares ocultas W..AN
for i in range(23,41):
    L=CL(i); rg.column_dimensions[L].width=11; rg.column_dimensions[L].hidden=True

# encabezados
rg['H18']='Dependencia'
rg['O18']='Ingrese\nDependencia'
rg['P18']='Tipo\nPart'
rg['S18']='Teléfono\nadicional'
rg['T18']='Correo adicional'
rg['E11']=('Ingrese la fecha y el ID o C.C. en las casillas sin relleno (título azul claro). '
           'Si el participante aparece como INEXISTENTE, complete las columnas L a O. '
           'Si no tiene el documento, escriba el nombre en el buscador.')

# mayúsculas en los tipos escritos a mano
for r in range(FIRST,LAST+1):
    v=rg['M%d'%r].value
    if isinstance(v,str) and v.strip(): rg['M%d'%r]=v.strip().upper()

# ---- fórmulas de la tabla ----
# W _doc | X _fCol | Y _fEst | Z _clave | AA _primera | AB _correo2
# AC/AD/AE estudiantes | AF/AG/AH profesores | AI/AJ/AK administrativos
for r in range(FIRST,LAST+1):
    f=lambda s: s.format(r=r,f=FIRST,l=LAST,o=FIRST-1)
    rg['A%d'%r]=f('=IF(OR($W{r}="",$E{r}=""),"",1)')
    rg['B%d'%r]=f('=IF(ISNUMBER($E{r}),DAY($E{r}),"")')
    rg['C%d'%r]=f('=IF(ISNUMBER($E{r}),MONTH($E{r}),"")')
    rg['W%d'%r]=f('=IF(TRIM(SUBSTITUTE($F{r}&"",CHAR(160),""))="","",'
                  'IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM($F{r}&""),'
                  'CHAR(160),"")," ",""),".",""),",",""),TRIM(SUBSTITUTE($F{r}&"",CHAR(160),""))))')
    rg['X%d'%r]=f('=IF($W{r}="","",IFERROR(MATCH($W{r},BD_COL_CC,0),IFERROR(MATCH($W{r},BD_COL_ID,0),"")))')
    rg['Y%d'%r]=f('=IF(OR($W{r}="",$X{r}<>""),"",IFERROR(MATCH($W{r},BD_EST_DOC,0),'
                  'IFERROR(MATCH($W{r},BD_EST_ID,0),"")))')
    rg['Z%d'%r]=f('=IF($W{r}="","",$W{r}&"")')
    rg['AA%d'%r]=f('=IF($Z{r}="","",IF(MATCH($Z{r},$Z${f}:$Z${l},0)=ROW()-{o},1,0))')
    rg['AB%d'%r]=f('=IF($W{r}="","",IF($Y{r}<>"",IFERROR(LEFT(INDEX(BD_EST_MAIL2,$Y{r})&"",'
                   'FIND("#",INDEX(BD_EST_MAIL2,$Y{r})&"")-1),INDEX(BD_EST_MAIL2,$Y{r})&""),'
                   'IF($X{r}<>"",INDEX(BD_COL_MAIL2,$X{r})&"","")))')
    rg['G%d'%r]=f('=IF($W{r}="","",UPPER(IF($Y{r}<>"",INDEX(BD_EST_SEDE,$Y{r})&"",'
                  'IF($X{r}<>"",INDEX(BD_COL_SEDE,$X{r})&"",TRIM($N{r}&"")))))')
    rg['H%d'%r]=f('=IF($W{r}="","",UPPER(IF($Y{r}<>"",INDEX(BD_EST_PROG,$Y{r})&"",'
                  'IF($X{r}<>"",INDEX(BD_COL_PROG,$X{r})&"",TRIM($O{r}&"")))))')
    rg['I%d'%r]=f('=IF($W{r}="","",IF($Y{r}<>"",INDEX(BD_EST_ID,$Y{r}),'
                  'IF($X{r}<>"",INDEX(BD_COL_ID,$X{r}),"")))')
    rg['J%d'%r]=f('=IF($W{r}="","",IF($Y{r}<>"",INDEX(BD_EST_NOM,$Y{r})&"",'
                  'IF($X{r}<>"",INDEX(BD_COL_NOM,$X{r})&"","INEXISTENTE")))')
    rg['K%d'%r]=f('=IF($W{r}="","",IF($Y{r}<>"",'
                  'IF(INDEX(BD_EST_MAIL,$Y{r})&""<>"",INDEX(BD_EST_MAIL,$Y{r})&"",TRIM($L{r}&"")),'
                  'IF($X{r}<>"",IF(INDEX(BD_COL_MAIL,$X{r})&""<>"",INDEX(BD_COL_MAIL,$X{r})&"",'
                  'TRIM($L{r}&"")),TRIM($L{r}&""))))')
    rg['P%d'%r]=f('=IF($W{r}="","",IF($Y{r}<>"",IF(INDEX(BD_EST_COD,$Y{r})=1,"ESTUDIANTE",""),'
                  'IF($X{r}<>"",IF(INDEX(BD_COL_COD,$X{r})=27,"PROFESOR",'
                  'IF(INDEX(BD_COL_COD,$X{r})=28,"ADMINISTRATIVO","")),UPPER(TRIM($M{r}&"")))))')
    rg['R%d'%r]=f('=IF($W{r}="","",IF($Y{r}<>"",INDEX(BD_EST_CEL,$Y{r})&"",'
                  'IF($X{r}<>"",INDEX(BD_COL_TEL,$X{r})&"","")))')
    rg['S%d'%r]=f('=IF(OR($W{r}="",$Y{r}=""),"",'
                  'IF(AND(INDEX(BD_EST_TEL2,$Y{r})&""<>"",INDEX(BD_EST_TEL2,$Y{r})&""<>$R{r}),'
                  'INDEX(BD_EST_TEL2,$Y{r})&"",'
                  'IF(AND(INDEX(BD_EST_TEL3,$Y{r})&""<>"",INDEX(BD_EST_TEL3,$Y{r})&""<>$R{r}),'
                  'INDEX(BD_EST_TEL3,$Y{r})&"","")))')
    rg['T%d'%r]=f('=IF(OR($AB{r}="",$AB{r}=$K{r}),"",$AB{r})')
    for tp,(kc,pc,ic) in (('ESTUDIANTE',('AC','AD','AE')),('PROFESOR',('AF','AG','AH')),
                          ('ADMINISTRATIVO',('AI','AJ','AK'))):
        rg['%s%d'%(kc,r)]='=IF(AND($P{r}="{t}",$H{r}<>""),$H{r}&"","")'.format(r=r,t=tp)
        rg['%s%d'%(pc,r)]=('=IF(${k}{r}="","",IF(MATCH(${k}{r},${k}${f}:${k}${l},0)=ROW()-{o},${k}{r},""))'
                           ).format(r=r,k=kc,f=FIRST,l=LAST,o=FIRST-1)
        rg['%s%d'%(ic,r)]='=IF(${p}{r}="","",COUNTIF(${p}${f}:${p}{r},"?*"))'.format(r=r,p=pc,f=FIRST)
print('fórmulas de Registros listas (%.0fs)'%(time.time()-t0))

# ---- buscador (mismo lugar y aspecto del original, sin tabla dinámica) ----
rg['G13']='Dependencia'; rg['J13']='Escriba el nombre'; rg['L13']='Escriba el correo'
for ref in ('J14','L14','J15'):
    c=rg[ref]; c.value=None; cp(c, rg['E14'])
    c.fill=INPUT_FILL; c.alignment=Alignment(horizontal='left',vertical='center')
    c.font=Font(name='Arial',size=10,bold=True,color='FF000000')
rg['K14']='← escriba el nombre   ·   o el correo →'
rg['AM14']=('=IF(TRIM($J$14&"")<>"",IFERROR(MATCH(TRIM($J$14),BD_EST_NOM,0),'
            'IFERROR(MATCH("*"&TRIM($J$14)&"*",BD_EST_NOM,0),"")),'
            'IF(TRIM($L$14&"")<>"",IFERROR(MATCH(TRIM($L$14),BD_EST_MAIL,0),'
            'IFERROR(MATCH("*"&TRIM($L$14)&"*",BD_EST_MAIL,0),"")),""))')
rg['AM15']=('=IF(TRIM($J$15&"")="","",IFERROR(MATCH(TRIM($J$15),BD_COL_NOM,0),'
            'IFERROR(MATCH("*"&TRIM($J$15)&"*",BD_COL_NOM,0),"")))')
for r,escrito,sede,idc,prog,doc in ((14,'$J$14&$L$14','BD_EST_SEDE','BD_EST_ID','BD_EST_PROG','BD_EST_DOC'),
                                    (15,'$J$15','BD_COL_SEDE','BD_COL_ID','BD_COL_PROG','BD_COL_CC')):
    p='$AM$%d'%r
    rg['E%d'%r]=('=IF(TRIM({e})="","",IF({p}="","No encontrado",INDEX({s},{p})&""))'
                 ).format(e=escrito,p=p,s=sede)
    rg['F%d'%r]='=IF({p}="","",INDEX({g},{p}))'.format(p=p,g=idc)
    rg['G%d'%r]='=IF({p}="","",INDEX({g},{p})&"")'.format(p=p,g=prog)
    rg['H%d'%r]='=IF({p}="","",INDEX({g},{p}))'.format(p=p,g=doc)
    rg['F%d'%r].number_format='0'; rg['H%d'%r].number_format='0'

# ---- cuadro de participación (espacio libre arriba a la derecha) ----
ST_PCT='0.0%'
for c in 'LMNOP': cp(rg[c+'1'], rg['E12'])
rg['L1']='RESUMEN DE PARTICIPACIÓN'
rg['L2']='Filtrar por sede:'
rg['L2'].font=Font(name='Arial',size=10,bold=True); rg['L2'].alignment=Alignment(horizontal='right')
cp(rg['M2'], rg['E14']); rg['M2'].fill=INPUT_FILL
rg['M2'].font=Font(name='Arial',size=10,bold=True); rg['M2']='TODAS'
for c,t in zip('LMNOP',['Tipo de participante','Ppantes','%','Ppaciones','%']):
    cp(rg[c+'3'], rg['E13']); rg[c+'3']=t
for k,t in enumerate(TIPOS):
    r=4+k
    for c in 'LMNOP': cp(rg['%s%d'%(c,r)], rg['E14'])
    rg['L%d'%r]=t; rg['L%d'%r].alignment=Alignment(horizontal='left',vertical='center')
    rg['M%d'%r]=('=IF($M$2="TODAS",COUNTIFS(REG_PRIM,1,REG_TIPO,"{t}"),'
                 'COUNTIFS(REG_PRIM,1,REG_TIPO,"{t}",REG_SEDE,$M$2))').format(t=t)
    rg['N%d'%r]='=IFERROR(M{r}/$M$8,"")'.format(r=r)
    rg['O%d'%r]=('=IF($M$2="TODAS",COUNTIFS(REG_QPART,1,REG_TIPO,"{t}"),'
                 'COUNTIFS(REG_QPART,1,REG_TIPO,"{t}",REG_SEDE,$M$2))').format(t=t)
    rg['P%d'%r]='=IFERROR(O{r}/$O$8,"")'.format(r=r)
for c in 'LMNOP': cp(rg[c+'8'], rg['E13'])
rg['L8']='TOTAL'; rg['L8'].alignment=Alignment(horizontal='left',vertical='center')
rg['M8']='=SUM(M4:M7)'; rg['N8']='=IF($M$8=0,"",SUM(N4:N7))'
rg['O8']='=SUM(O4:O7)'; rg['P8']='=IF($O$8=0,"",SUM(P4:P7))'
for r in range(4,9):
    rg['N%d'%r].number_format=ST_PCT; rg['P%d'%r].number_format=ST_PCT
    rg['M%d'%r].number_format='#,##0'; rg['O%d'%r].number_format='#,##0'

# ---- validaciones ----
from openpyxl.worksheet.datavalidation import DataValidationList
rg.data_validations=DataValidationList()
v=dv(rg,'DATE(2015,1,1)',['E%d:E%d'%(FIRST,LAST)],tipo='date',operator='greaterThan')
v.errorStyle='warning'; v.showErrorMessage=True; v.showInputMessage=True
v.promptTitle='Ingreso de fecha:'; v.prompt='Ingrese fecha DD/MM/AAAA, con números. Ejemplo: 10/06/2020'
dv(rg,'"%s"'%','.join(TIPOS),['M%d:M%d'%(FIRST,LAST)])
dv(rg,'"%s"'%','.join(SEDES),['N%d:N%d'%(FIRST,LAST)])
dv(rg,'"TODAS,%s"'%','.join(SEDES),['M2'])

# ---- formato condicional ----
rg.conditional_formatting=ConditionalFormattingList()
rg.conditional_formatting.add('I%d:I%d'%(FIRST,LAST),
    Rule(type='duplicateValues',dxf=DifferentialStyle(font=Font(bold=True,color='FF0000FF'))))
rg.conditional_formatting.add('I%d:I%d'%(FIRST,LAST),
    CellIsRule(operator='equal',formula=['0'],font=Font(color='FFFFFFFF')))
rg.conditional_formatting.add('E%d:T%d'%(FIRST,LAST),
    FormulaRule(formula=['$J%d="INEXISTENTE"'%FIRST],fill=PatternFill('solid',fgColor='FFFFD9CC')))
rg.freeze_panes=None
rg.sheet_view.topLeftCell='A1'
rg.sheet_view.selection[0].activeCell='F19'; rg.sheet_view.selection[0].sqref='F19'

# ======================= RESULTADOS POR TIPO: se elimina =======================
del wb['Resultados por tipo']
print('Registros listo (%.0fs)'%(time.time()-t0))

# ======================= RESULTADOS POR PROGRAMA (layout original) =======================
rp=wb['Resultados por programa']
rp['B3']='Sede (se elige en Registros):'
rp['D3']='=Registros!$M$2'
rp['F5']='PROFESORES PARTICIPANTES'
for c in ('B7','F7','J7'): rp[c]='Dependencia'
H0,H1,TOT = 8,37,38
for i in range(14,26):                                  # N..Y auxiliares ocultas
    L=CL(i); rp.column_dimensions[L].width=10; rp.column_dimensions[L].hidden=True
for r in range(H0,TOT+1):
    for c in range(14,26): rp.cell(row=r,column=c).value=None
def bloque(cini,tipo,pn,ix,aux):
    c0,c1,c2=CL(cini),CL(cini+1),CL(cini+2)
    ap,ac,ak,ao=aux
    cnt=lambda crit,ref:('IF($D$3="TODAS",COUNTIFS({c},1,REG_TIPO,"{t}",REG_PROG,{r}),'
                         'COUNTIFS({c},1,REG_TIPO,"{t}",REG_PROG,{r},REG_SEDE,$D$3))').format(c=crit,t=tipo,r=ref)
    for r in range(H0,H1+1):
        rp['%s%d'%(ap,r)]=('=IFERROR(INDEX({pn},MATCH(ROWS(${a}${h}:{a}{r}),{ix},0))&"","")'
                           ).format(pn=pn,ix=ix,a=ap,h=H0,r=r)
        rp['%s%d'%(ac,r)]='=IF(${a}{r}="","",{f})'.format(a=ap,r=r,f=cnt('REG_PRIM','$%s%d'%(ap,r)))
        rp['%s%d'%(ak,r)]='=IF(${a}{r}="","",${c}{r}*10000+(10000-ROW()))'.format(a=ap,c=ac,r=r)
        rp['%s%d'%(ao,r)]=('=IFERROR(MATCH(LARGE(${k}${h}:${k}${e},ROWS(${o}${h}:{o}{r})),'
                           '${k}${h}:${k}${e},0),"")').format(k=ak,o=ao,h=H0,e=H1,r=r)
        rp['%s%d'%(c0,r)]='=IF(${o}{r}="","",INDEX(${a}${h}:${a}${e},${o}{r})&"")'.format(o=ao,a=ap,h=H0,e=H1,r=r)
        rp['%s%d'%(c1,r)]='=IF(${o}{r}="","",INDEX(${c}${h}:${c}${e},${o}{r}))'.format(o=ao,c=ac,h=H0,e=H1,r=r)
        rp['%s%d'%(c2,r)]='=IF(${c0}{r}="","",{f})'.format(c0=c0,r=r,f=cnt('REG_QPART','$%s%d'%(c0,r)))
    rp['%s%d'%(c0,TOT)]='Total'
    rp['%s%d'%(c1,TOT)]='=SUM(%s%d:%s%d)'%(c1,H0,c1,H1)
    rp['%s%d'%(c2,TOT)]='=SUM(%s%d:%s%d)'%(c2,H0,c2,H1)
bloque(2,'ESTUDIANTE','REG_PROGEST','REG_IEST',('N','O','P','Q'))
bloque(6,'PROFESOR','REG_PROGPRO','REG_IPRO',('R','S','T','U'))
bloque(10,'ADMINISTRATIVO','REG_PROGADM','REG_IADM',('V','W','X','Y'))

# ============================ EVALUACIÓN (layout original) ============================
ev=wb['Evaluación']
ev['A2']='Formato FR-BM-DFB-03  ·  Versión 1  ·  Enero 28 de 2021'
ev['A2'].font=Font(name='Arial',size=8,italic=True,color='FF595959')
ev['O3']=('E - Excelente (4)    N - Notable (3)    A - Aceptable (2)    '
          'N/M - Necesita Mejoramiento (1)    N/A - No aplica')
ev['O3'].alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
ev['B6']='Ingrese resultados de evaluaciones (se puede copiar y pegar)'
ENT=['B','C','D','E','F','G','H','I','J','K']
ESP=['P','Q','R','T','U','V','W','X','Z','AA']
E0,E1=10,509
for s,d in zip(ENT,ESP):
    ev['%s5'%d]='=IFERROR(AVERAGEIF({d}{a}:{d}{b},">0"),"")'.format(d=d,a=E0,b=E1)
    ev['%s6'%d]='=IF({d}5<>"",{d}5*0.25,"")'.format(d=d)
ev['AB6']='=IFERROR(AVERAGE(P6:R6,T6:X6,Z6:AA6),"")'
for r in range(E0,E1+1):
    for s,d in zip(ENT,ESP):
        t='TRIM(${s}{r}&"")'.format(s=s,r=r); u='UPPER(%s)'%t
        ev['%s%d'%(d,r)]=(
            '=IF({t}="","",IFERROR(IF(AND(VALUE({t})>=0,VALUE({t})<=4),VALUE({t}),"-"),'
            'IF({u}="E",4,IF({u}="N",3,IF({u}="A",2,'
            'IF(OR({u}="N/M",{u}="NM"),1,IF(OR({u}="N/A",{u}="NA"),0,"-")))))))').format(t=t,u=u)
from openpyxl.worksheet.datavalidation import DataValidationList
ev.data_validations=DataValidationList()
v=dv(ev,'"E,N,A,N/M,N/A,4,3,2,1,0"',['B%d:K%d'%(E0,E1)]); v.errorStyle='warning'; v.showErrorMessage=True
ev.conditional_formatting=ConditionalFormattingList()
for s,d in zip(ENT,ESP):
    ev.conditional_formatting.add('{s}{a}:{s}{b}'.format(s=s,a=E0,b=E1),
        FormulaRule(formula=['${d}{a}="-"'.format(d=d,a=E0)],
                    fill=PatternFill('solid',fgColor='FFFFC7CE'),font=Font(bold=True,color='FF9C0006')))
print('Programa y Evaluación listos (%.0fs)'%(time.time()-t0))

# ============================ GRÁFICOS EVALUACIÓN ============================
# el original tenía 4 gráficos que usted encontró desordenados: se deja uno solo
idx=wb.sheetnames.index('Gráficos Evaluación')
del wb['Gráficos Evaluación']
gr=wb.create_sheet('Gráficos Evaluación', idx)
gr.sheet_view.showGridLines=False
for k,v in {'A':2,'B':5,'C':38,'D':11}.items(): gr.column_dimensions[k].width=v
AZ='FF1F3864'; AZ2='FF2E5B9A'; GRIS='FFD9D9D9'
def A(sz=10,b=False,c='FF000000',i=False): return Font(name='Arial',size=sz,bold=b,color=c,italic=i)
gr['B1']='GRÁFICAS DE EVALUACIÓN'; gr['B1'].font=A(14,True)
PREG=[('Divulgación de la actividad','P6','1. La divulgación de la actividad, evento o servicio fue'),
 ('Cumplimiento de los tiempos','Q6','2. La actividad, evento o servicio se realizó de acuerdo con los tiempos establecidos'),
 ('Recursos tecnológicos y audiovisuales','R6','3. Los recursos tecnológicos y audiovisuales utilizados fueron'),
 ('Claridad de la temática','T6','4. La claridad en la temática o propósito de la actividad, evento o servicio fue'),
 ('Participación de los asistentes','U6','5. La actividad, evento o servicio promovió la participación activa de los asistentes'),
 ('Actitud del facilitador','V6','6. La Actitud y disponibilidad del facilitador o colaborador fue'),
 ('Manejo del tema','W6','7. El manejo del tema por parte del facilitador o colaborador fue'),
 ('Lugar y condiciones físicas','X6','8. El lugar y sus condiciones físicas fueron'),
 ('Cumplimiento de expectativas','Z6','9. Se cumplió su expectativa frente a la actividad, evento o servicio'),
 ('Aporte a la formación integral','AA6','10. La actividad, evento o servicio realizado aportó a su formación integral')]
for c,t in zip('BCD',['Nº','Aspecto evaluado','%']):
    x=gr[c+'3']; x.value=t; x.font=A(10,True,'FFFFFFFF'); x.fill=PatternFill('solid',fgColor=AZ2)
    x.alignment=Alignment(horizontal='center',vertical='center'); x.border=BOX
for k,(corta,src,larga) in enumerate(PREG):
    r=4+k
    gr['B%d'%r]=k+1; gr['C%d'%r]=corta; gr['D%d'%r]='=Evaluación!%s'%src
    for c in 'BCD':
        x=gr['%s%d'%(c,r)]; x.font=A(10); x.border=BOX
        x.alignment=Alignment(horizontal='left' if c=='C' else 'center',vertical='center')
    gr['D%d'%r].number_format='0.0%'; gr['D%d'%r].font=A(10,True,AZ)
    gr['D%d'%r].fill=PatternFill('solid',fgColor='FFF2F2F2')
for c,t in zip('BCD',['','PROMEDIO TOTAL','=Evaluación!AB6']):
    x=gr[c+'15']; x.value=t if t else None; x.font=A(11,True,'FFFFFFFF')
    x.fill=PatternFill('solid',fgColor=AZ); x.border=BOX; x.alignment=Alignment(horizontal='center')
gr['D15'].number_format='0.0%'
ch=BarChart(); ch.type='bar'; ch.title=None; ch.legend=None
ch.y_axis.numFmt='0%'; ch.y_axis.scaling.min=0; ch.y_axis.scaling.max=1
ch.y_axis.majorGridlines=None; ch.x_axis.majorGridlines=None
ch.x_axis.delete=False; ch.y_axis.delete=False
ch.x_axis.scaling.orientation='maxMin'            # la pregunta 1 arriba
ch.gapWidth=45
ch.add_data(Reference(gr,min_col=4,min_row=4,max_row=13),titles_from_data=False)
ch.series[0].cat=AxDataSource(strRef=StrRef(f="'Gráficos Evaluación'!$C$4:$C$13"))
ch.series[0].graphicalProperties=GraphicalProperties(solidFill='2E5B9A')
dl=DataLabelList(); dl.showVal=True; dl.showSerName=False; dl.showCatName=False
dl.showLegendKey=False; dl.showPercent=False; dl.showBubbleSize=False
ch.dataLabels=dl
ch.width=19; ch.height=11
gr.add_chart(ch,'F3')
gr['B18']='Preguntas completas (formato FR-BM-DFB-03)'; gr['B18'].font=A(9,True,'FF595959')
for k,(corta,src,larga) in enumerate(PREG):
    gr['B%d'%(19+k)]=larga; gr['B%d'%(19+k)].font=A(8,False,'FF808080')

# ============================ NOMBRES DEFINIDOS ============================
EST=lambda c:"'BD EST'!${c}:${c}".format(c=c)
COL=lambda c:"'BD ADM-DOC'!${c}:${c}".format(c=c)
RG =lambda c:"Registros!${c}${f}:${c}${l}".format(c=c,f=FIRST,l=LAST)
NAMES={
 'BD_EST_ID':EST('B'),'BD_EST_NOM':EST('C'),'BD_EST_SEDE':EST('G'),'BD_EST_PROG':EST('K'),
 'BD_EST_CEL':EST('V'),'BD_EST_TEL2':EST('W'),'BD_EST_TEL3':EST('X'),
 'BD_EST_MAIL':EST('Y'),'BD_EST_MAIL2':EST('Z'),'BD_EST_DOC':EST('AB'),'BD_EST_COD':EST('AU'),
 'BD_COL_CC':COL('B'),'BD_COL_ID':COL('C'),'BD_COL_NOM':COL('D'),'BD_COL_COD':COL('I'),
 'BD_COL_PROG':COL('K'),'BD_COL_MAIL':COL('L'),'BD_COL_MAIL2':COL('M'),'BD_COL_TEL':COL('N'),
 'BD_COL_SEDE':COL('P'),
 'REG_QPART':RG('A'),'REG_SEDE':RG('G'),'REG_PROG':RG('H'),'REG_TIPO':RG('P'),'REG_PRIM':RG('AA'),
 'REG_PROGEST':RG('AD'),'REG_IEST':RG('AE'),'REG_PROGPRO':RG('AG'),'REG_IPRO':RG('AH'),
 'REG_PROGADM':RG('AJ'),'REG_IADM':RG('AK')}
for n2,v2 in NAMES.items(): wb.defined_names.add(DefinedName(n2,attr_text=v2))

wb.active=wb.sheetnames.index('Registros')
for ws in wb.worksheets: ws.sheet_view.tabSelected = (ws.title=='Registros')
wb.calculation.fullCalcOnLoad=True
raw=SP+'/_v17_raw.xlsx'
wb.save(raw)
from sharedstr import trim_empty_cells
trim=SP+'/_v17_trim.xlsx'
print('limpiadas:',trim_empty_cells(raw,trim,{'BD EST','BD ADM-DOC'}))
out=SP+'/Registro_Asistencia_2026-1_V17.xlsx'
n,size=convert(trim,out)
print('cadenas compartidas:',n,'| tamaño: %.1f MB'%(size/1048576))
print('guardado:',out,'(%.0fs)'%(time.time()-t0))
